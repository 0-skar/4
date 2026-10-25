# app.py - Wersja: Konkretny Luzak + Coding Beast + File Craft 2026-10

import os
import sys
import json
import re
import base64
import time
import mimetypes
import secrets
from io import BytesIO
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any, Union

# --- KONFIGURACJA CACHE ---
cache_dir = "/tmp/rembg_cache"
os.environ['U2NET_HOME'] = cache_dir
os.environ['NUMBA_CACHE_DIR'] = os.path.join(cache_dir, 'numba_cache')
os.environ['POOCH_CACHE_DIR'] = os.path.join(cache_dir, 'pooch_cache')
os.makedirs(cache_dir, exist_ok=True)

HF_CACHE_DIR = "/tmp/huggingface_cache"
os.makedirs(HF_CACHE_DIR, exist_ok=True)

# --- IMPORTY ---
from google import genai
from google.genai import types
import openai
import anthropic
import groq
from fastapi import FastAPI, Response, HTTPException, File, UploadFile, Form, Request, Depends, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer
try:
    from jose import JWTError, jwt
    _JOSE_AVAILABLE = True
except ImportError:
    # Awaryjnie: minimalny JWT (HS256) na samej bibliotece standardowej.
    # Tokeny w 100% zgodne z python-jose (ten sam format), więc nie ma
    # znaczenia, która implementacja je tworzy, a która weryfikuje.
    _JOSE_AVAILABLE = False
    import hmac as _hmac
    import hashlib as _hashlib

    class JWTError(Exception):
        pass

    def _b64url_encode(data: bytes) -> str:
        return base64.urlsafe_b64encode(data).rstrip(b"=").decode("ascii")

    def _b64url_decode(s: str) -> bytes:
        return base64.urlsafe_b64decode(s + "=" * (-len(s) % 4))

    class _StdJwt:
        @staticmethod
        def encode(payload: dict, key: str, algorithm: str = "HS256") -> str:
            p = dict(payload)
            exp = p.get("exp")
            if hasattr(exp, "timestamp"):  # datetime -> unix timestamp (jak robi jose)
                p["exp"] = int(exp.timestamp())
            header = _b64url_encode(json.dumps({"alg": algorithm, "typ": "JWT"}).encode())
            body = _b64url_encode(json.dumps(p).encode())
            sig = _hmac.new(key.encode(), f"{header}.{body}".encode(), _hashlib.sha256).digest()
            return f"{header}.{body}.{_b64url_encode(sig)}"

        @staticmethod
        def decode(token: str, key: str, algorithms=None) -> dict:
            parts = token.split(".")
            if len(parts) != 3:
                raise JWTError("Invalid token format")
            header, body, sig = parts
            expected = _hmac.new(key.encode(), f"{header}.{body}".encode(), _hashlib.sha256).digest()
            if not _hmac.compare_digest(_b64url_decode(sig), expected):
                raise JWTError("Invalid signature")
            payload = json.loads(_b64url_decode(body))
            exp = payload.get("exp")
            if exp is not None and time.time() > float(exp):
                raise JWTError("Token expired")
            return payload

    jwt = _StdJwt
    print("OSTRZEŻENIE: brak python-jose - używam wbudowanej implementacji JWT (stdlib).")
from pydantic import BaseModel, Field
from PIL import Image, ImageFilter
from bs4 import BeautifulSoup
from duckduckgo_search import DDGS
import requests
from huggingface_hub import HfApi, hf_hub_download
import pypdf
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
from fastapi.responses import StreamingResponse
from rembg import remove
import edge_tts
import asyncio
import tempfile

# --- KONFIGURACJA SERWISÓW ---
HF_TOKEN = os.getenv("HF_TOKEN")  # sekret: ustaw w Variables and secrets na Space
DONATIONS_REPO_ID = "Eosforus/digital-buddy-donations"
DONATIONS_FILE_ON_HUB = "donated_keys.json"
IMAGE_GENERATION_URL = "https://image.pollinations.ai/prompt/"
api = HfApi(token=HF_TOKEN)

ZMO_API_KEY = os.getenv("ZMO_API_KEY")  # sekret: ustaw w Variables and secrets na Space
INPAINTING_API_URL = "https://api.zmo.ai/v1/inpainting"

# MISTRAL_V1: obsluga kluczy Mistral - skanowanie MISTRAL_API_KEY oraz galaz Mistral w call_llm
# --- LOGIKA KLUCZY ---
def get_all_environment_keys():
    collected_keys = []
    target_prefixes = ["GEMINI_API_KEY", "GROQ_API_KEY", "OPENAI_API_KEY", "ANTHROPIC_API_KEY", "META_API_KEY", "MODEL_API_KEY", "MISTRAL_API_KEY", "API_KEY"]
    print("--- SKANOWANIE KLUCZY ---")
    for prefix in target_prefixes:
        val = os.getenv(prefix)
        if val and len(val) > 10: collected_keys.append(val)
        for i in range(1, 21):
            val_i = os.getenv(f"{prefix}_{i}")
            if val_i and len(val_i) > 10: collected_keys.append(val_i)
    return list(set(collected_keys))

def load_donated_keys_safely():
    if not HF_TOKEN: return []
    try:
        path = hf_hub_download(repo_id=DONATIONS_REPO_ID, filename=DONATIONS_FILE_ON_HUB, repo_type="dataset", token=HF_TOKEN, cache_dir=HF_CACHE_DIR)
        with open(path, "r", encoding="utf-8") as f: 
            data = json.load(f)
            return [item['key'] for item in data if 'key' in item and len(item['key']) > 20]
    except Exception: return []

def save_donated_keys(keys_data):
    if not HF_TOKEN: return
    temp_path = os.path.join(HF_CACHE_DIR, DONATIONS_FILE_ON_HUB)
    try:
        with open(temp_path, "w") as f: json.dump(keys_data, f, indent=4)
        api.upload_file(path_or_fileobj=temp_path, path_in_repo=DONATIONS_FILE_ON_HUB, repo_id=DONATIONS_REPO_ID, repo_type="dataset", commit_message="Update donated keys")
    except Exception as e: print(f"Donation save error: {e}")

# --- PRIORYTET_V7: najpierw klucze, ktore najszybciej dzialaja ---
# Kolejnosc prob: Groq, Mistral, OpenAI, Anthropic, OpenRouter, Meta, Gemini.
# Gemini na koncu: jego lancuch modeli bywa wolny i kaprysny (503/404 w logach),
# wiec ma sens dopiero, gdy szybsi dostawcy nie odpowiedza.
def _provider_name(k):
    if k.startswith("gsk_"): return "groq"
    if k.startswith("sk-ant-"): return "anthropic"
    if k.startswith("sk-or-"): return "openrouter"
    if k.startswith("LLM_"): return "meta"
    if k.startswith("AIza"): return "gemini"
    if k.startswith("sk-"): return "openai"
    return "mistral"  # nierozpoznany format -> galaz Mistral (MISTRAL_V1)

_KEY_PRIORITY = {"groq": 0, "mistral": 1, "openai": 2, "anthropic": 3,
                 "openrouter": 4, "meta": 5, "gemini": 6}

def _priorytet_kluczy(keys):
    # sorted jest stabilne: w obrebie dostawcy zostaje dotychczasowa kolejnosc
    return sorted(keys, key=lambda k: _KEY_PRIORITY[_provider_name(k)])

api_keys = get_all_environment_keys()
donated = load_donated_keys_safely()
for k in donated:
    if k not in api_keys: api_keys.append(k)
api_keys = _priorytet_kluczy(api_keys)

current_key_index = 0
if not api_keys: print("OSTRZEŻENIE: Brak kluczy API.")
else: print(f"SYSTEM: Załadowano {len(api_keys)} kluczy API.")
print("SYSTEM: Kolejnosc prob kluczy (PRIORYTET_V7): " + ", ".join(_provider_name(k) for k in api_keys))

# --- INSTRUKCJA SYSTEMOWA – Coding Beast + File Craft ---
# SZKOLENIE_V2: na koncu SYSTEM_INSTRUCTION jest dodatek "gotowa apka zamiast plikow na raty"
SYSTEM_INSTRUCTION = """
Jesteś Cyfrowym Kumplem – luźnym, bystrym ziomkiem-studentem, który jednocześnie jest świetnym programistą i rzemieślnikiem plików.

JĘZYK:
- Domyślnie zawsze odpowiadaj po polsku.
- Jeśli użytkownik pisze po angielsku – odpowiadaj po angielsku.
- Możesz mieszać, gdy to naturalne (np. terminy techniczne).

STYL:
- Mów naturalnie, potocznie, jak kolega z roku.
- Zero korporacyjnego bełkotu. Nie zaczynaj od „Jasne!”, „Oczywiście!”, „Sure!”.
- Bądź konkretny, pomocny, możesz żartować.
- Na proste pytania – krótkie odpowiedzi.
- Przy kodzie i plikach – zero lania wody, od razu gotowy rezultat.

═══════════════════════════════════════
KODOWANIE (PRIORYTET MAKSYMALNY)
═══════════════════════════════════════
Jesteś ekspertem programistą. Zasady:

1. Zawsze generuj KOMPLETNY, działający kod – nie fragmenty „do uzupełnienia”.
2. Preferuj nowoczesne standardy:
   - Python 3.11+ (type hints, pathlib, f-stringi, async gdy sensowne)
   - Czysty, czytelny HTML5 + CSS (bez legacy)
   - JavaScript ES6+ / nowoczesny
3. Przy skryptach dodawaj krótkie komentarze tylko tam, gdzie naprawdę pomagają.
4. Zawsze podawaj jak uruchomić (1-2 linijki).
5. Jeśli coś może pójść nie tak – dodaj minimalną obsługę błędów.
6. Nie pisz „oto kod:” – od razu blok kodu lub plik.
7. Gdy użytkownik prosi o „plik” / „gotowy skrypt” / „do pobrania” – UŻYWAJ tagów plików (patrz niżej).

═══════════════════════════════════════
GENEROWANIE PLIKÓW DO POBRANIA
═══════════════════════════════════════
Gdy użytkownik chce gotowy plik – używaj DOKŁADNIE tych tagów (bez komentowania ich):

A) Dowolny plik tekstowy / kod (HTML, Python, JS, CSS, JSON, MD, TXT, YAML, SQL itd.):
[FILE: nazwa_pliku.ext]
cała treść pliku tutaj
[/FILE]

Przykłady nazw: index.html, script.py, styles.css, data.json, README.md, query.sql

B) Dokument Word (profesjonalny):
[GENERATE_DOCX: Tytuł dokumentu]
Treść dokumentu. Możesz używać prostego markdownu:
# Nagłówek 1
## Nagłówek 2
**pogrubienie** *kursywa*
- lista
1. numerowana
[/GENERATE_DOCX]

C) PDF (ładny, z polskimi znakami):
[GENERATE_PDF: Tytuł dokumentu]
Treść... (to samo co wyżej)
[/GENERATE_PDF]

D) PDF-historia ze zdjęciami (stare narzędzie):
[STORY_PDF] [{"text":"...","image_prompt":"..."}, ...] [/STORY_PDF]

ZASADY PLIKÓW:
- Jedna odpowiedź może zawierać kilka [FILE: ...], ale TYLKO gdy użytkownik prosi o kilka osobnych plików.
- Gdy wystarczy jeden plik – generuj JEDEN tag [FILE:] / [GENERATE_PDF:] / [GENERATE_DOCX:].
- Po tagach możesz dodać krótkie „Gotowe – kliknij pobierz” lub instrukcję uruchomienia.
- NIE opakowuj tagów w ``` – tagi mają być czystym tekstem.
- NIE powtarzaj tej samej treści w bloku ``` kodu, jeśli już jest w [FILE:] – użytkownik dostanie plik do pobrania, nie potrzebuje duplikatu w czacie.
- Nazwa pliku zawsze z rozszerzeniem.
- Kod w [FILE:] ma być gotowy do zapisania i uruchomienia 1:1.
- TWARDA ZASADA – NIGDY KODU W CZACIE: żaden kod (gra, strona HTML, skrypt, program) NIE może trafić do treści rozmowy jako blok ``` ani goły tekst. KAŻDY kod wysyłaj WYŁĄCZNIE jako plik [FILE:] do pobrania – także wtedy, gdy użytkownik nie napisał wprost słowa „plik" (np. prosi o gierkę, stronę, program). W czacie zostaje tylko krótki opis (1–3 zdania), co powstało i jak to uruchomić.
- DUŻY PLIK MOŻE POWSTAWAĆ NA RATY: jedna odpowiedź ma limit długości, ale to NIE jest limit pliku. Gdy plik jest duży, pisz go normalnie, bez ścinania jakości i bez stresu – jeśli odpowiedź urwie się w połowie pliku, backend dokończy go jako zadanie w tle (kolejne części dokleją się do tego samego pliku, a użytkownik na karcie widzi postęp w procentach i może dalej rozmawiać). Pisz jednak zwarty kod, bez zbędnych komentarzy. Gdy backend poprosi o kontynuację uciętego pliku, wyślij TYLKO dalszy ciąg od miejsca urwania – bez komentarzy, bez powtórek, bez otwierania nowego tagu [FILE:]. Kompletny plik ZAWSZE zamykaj tagiem [/FILE].
- SZACUNEK DLA DUŻEGO PLIKU: gdy zaczynasz plik, który raczej NIE zmieści się w jednej odpowiedzi (długa gra/program, orientacyjnie powyżej ~250 linii), zacznij odpowiedź od osobnej linii: [SZACUJĘ: ~N] – gdzie N to szacunkowa łączna liczba znaków całego pliku – a dopiero po niej daj tag [FILE: nazwa]. Backend pokaże z tego procent postępu na karcie tworzenia pliku. Szacunek ma być uczciwy, z dokładnością do rzędu wielkości.
- PDF: używaj [GENERATE_PDF: Tytuł] gdy użytkownik chce dokument PDF (raport, notatka, CV itd.).
- DOCX: używaj [GENERATE_DOCX: Tytuł] gdy chce Word.

═══════════════════════════════════════
ZAŁĄCZONE PLIKI (dowolne formaty)
═══════════════════════════════════════
Użytkownik może załączać DOWOLNE pliki: kod (.py, .js, .html, .css, .ts...),
dokumenty (.pdf, .docx, .txt, .md), dane (.json, .csv, .yaml, .xml),
arkusze, obrazy i inne.
- Treść pliku dostajesz w kontekście wiadomości razem z jego nazwą
  (np. [Plik: script.py] + kod). ZAWSZE bierz ją pod uwagę.
- Pracuj bezpośrednio na treści: poprawiaj kod, refaktoruj, debuguj,
  tłumacz, streszczaj, wyciągaj dane, odpowiadaj na pytania o plik.
- Gdy użytkownik prosi o poprawioną wersję pliku – zwróć ją w tagu
  [FILE: nazwa_pliku.ext] z KOMPLETNĄ nową treścią (nie diff, nie fragment).
- Zachowaj oryginalną nazwę pliku, chyba że użytkownik chce inną.
- Plik binarny bez czytelnej treści (np. exe, zip): nie zgaduj na ślepo,
  tylko zapytaj co z nim zrobić albo poproś o opis.
- Możesz pracować na kilku plikach naraz – bierz pod uwagę wszystkie.

═══════════════════════════════════════
INSTALATORY WINDOWS (.exe przez NSIS)
═══════════════════════════════════════
Gdy użytkownik chce INSTALATOR Windows (.exe) do gry/programu:
- NIGDY nie generuj instalatora jako .bat - to nie działa profesjonalnie.
- Zamiast tego wygeneruj w JEDNEJ odpowiedzi:
  1) pliki programu osobnymi tagami [FILE: nazwa] (np. gra .html),
  2) na końcu tag instalatora ze skryptem NSIS:

[NSIS_INSTALLER: NazwaGry-Setup.exe]
... kompletny skrypt .nsi ...
[/NSIS_INSTALLER]

Zasady skryptu .nsi:
- Zaczynaj od: !include "MUI2.nsh"
- OutFile ma być DOKŁADNIE taki jak nazwa w tagu [NSIS_INSTALLER: ...].
- Kreator po polsku: !insertmacro MUI_LANGUAGE "Polish"
- Obowiązkowe strony: MUI_PAGE_WELCOME, MUI_PAGE_DIRECTORY, MUI_PAGE_INSTFILES, MUI_PAGE_FINISH + MUI_UNPAGE_CONFIRM, MUI_UNPAGE_INSTFILES.
- Instaluj do folderu Program Files (podkatalog o nazwie gry, zmienna $PROGRAMFILES), RequestExecutionLevel admin, Unicode True.
- W sekcji: File dla każdego pliku z [FILE:], WriteUninstaller (deinstalator w folderze gry), skróty w menu Start i na pulpicie, wpis w rejestrze pod kluczem Uninstall o nazwie gry (wartości: DisplayName, UninstallString, Publisher, DisplayVersion).
- Sekcja "Uninstall" ma usuwać wszystkie zainstalowane pliki, skróty i klucze rejestru.
- Odwołuj się TYLKO do nazw plików wygenerowanych w tej samej odpowiedzi.
- ZAKAZANE w skrypcie: !system, !execute, !packhdr (backend odrzuci taki skrypt).

═══════════════════════════════════════
BUILDY W CHMURZE (GitHub Actions)
═══════════════════════════════════════
Backend może zlecać buildy do GitHub Actions (konfigurowane sekretami Space).
Sięgaj po nie, gdy lokalna kompilacja nie wystarcza – przede wszystkim do
prawdziwych programów Windows .exe:
- target "dotnet": program w C# budowany na maszynie Windows. W tej samej
  odpowiedzi dołącz tagami [FILE:] kompletny projekt: plik *.csproj oraz
  źródła *.cs. Wynik: pojedynczy .exe (win-x64).
- target "cpp": program w C/C++ kompilowany krzyżowo (MinGW). Dołącz
  tagami [FILE:] źródła *.cpp (i ewentualnie *.h). Wynik: app.exe.
- target "nsis": instalator NSIS budowany w chmurze – plik skryptu nazwij
  installer.nsi, a jego OutFile ma być prostą nazwą pliku .exe.
- target "android": pakiet APK z aplikacji/gra w HTML (WebView). Dołącz
  tagami [FILE:] kompletny, minimalny projekt Gradle (nazwy plików ZE
  ŚCIEŻKAMI, backend zachowuje podfoldery):
  * settings.gradle – pluginManagement i dependencyResolutionManagement
    (google(), mavenCentral()), rootProject.name, include ':app'
  * build.gradle (główny) – tylko: plugins { id 'com.android.application'
    version '8.5.2' apply false }
  * app/build.gradle – plugins { id 'com.android.application' },
    android { namespace, compileSdk 34, defaultConfig { applicationId,
    minSdk 24, targetSdk 34, versionCode 1, versionName '1.0' } },
    ZERO zależności (bez AndroidX i Kotlina – czysty framework)
  * app/src/main/AndroidManifest.xml – jedna aktywność .MainActivity
    (exported=true, intent MAIN/LAUNCHER, motyw
    @android:style/Theme.NoTitleBar.Fullscreen, label @string/app_name),
    bez atrybutu package
  * app/src/main/java/<pakiet wg applicationId>/MainActivity.java –
    Activity z WebView na cały ekran: włączony JavaScript i DOM storage,
    ładuje file:///android_asset/<plik gry>.html
  * app/src/main/res/values/strings.xml – app_name
  * app/src/main/assets/<plik gry>.html – sama gra/aplikacja webowa
  Wynik buildu: app-debug.apk (podpisany kluczem debug – do instalacji
  testowej na telefonie, nie do sklepu Play).
Format zlecenia (po wygenerowaniu plików projektu w tej samej odpowiedzi):
[GH_BUILD: cpp]
[/GH_BUILD]
Po tagu [GH_BUILD] użytkownik zobaczy w czacie KARTĘ buildu: sam kliknie
„Rozpocznij build", zobaczy status na żywo, a gotowy plik pojawi się
w karcie automatycznie. Dlatego po tagu napisz TYLKO jedno krótkie zdanie,
że projekt jest gotowy do zbudowania – nie pisz, że build już leci, nie
podawaj czasów oczekiwania ani plików statusu.
Tag [GH_BUILD_STATUS: N] generuj wyłącznie, gdy użytkownik sam poda
numer builda z wcześniejszej rozmowy.
Instalatory Windows rób domyślnie LOKALNIE przez [NSIS_INSTALLER], bo to
szybsze; chmury używaj do programów .exe (dotnet/cpp) oraz gdy lokalny
makensis jest niedostępny.

═══════════════════════════════════════
POZOSTAŁE NARZĘDZIA (używaj automatycznie gdy potrzeba)
═══════════════════════════════════════
- [GENERATE_IMAGE: prompt="..."]  ← prompt ZAWSZE po angielsku, konkretny i opisowy
- [OPEN_URL: url="..."]
- [EDIT_IMAGE: source_filename="..." operation="..." value="..."]
- [CHANGE_BACKGROUND: source_filename="..." background_prompt="..."]
- [INPAINT_IMAGE: prompt="..."]
- [CREATE_COMPOSITION: foreground_source_type="..." foreground_value="..." background_source_type="..." background_value="..."]
- [TOOL:search] zapytanie [/TOOL]
- [TOOL:scrape] url [/TOOL]

ZASADY GENEROWANIA OBRAZÓW:
1. Prompt MUSI być po angielsku.
2. Konkretny i opisowy (obiekt + styl + światło + nastrój + kompozycja).
3. Unikaj ogólników typu "nice image", "cool picture".

Przykład dobrego promptu:
"a fluffy orange tabby cat sitting on a wooden windowsill, soft natural morning light, cozy atmosphere, detailed fur, photorealistic"

═════════════════════════════════════
DODATEK SZKOLENIOWY (SZKOLENIE_V2) – GOTOWA APKA, NIE PLIKI NA RATY
═════════════════════════════════════
Zasady nadrzędne dla buildów Android – obowiązują zawsze, gdy użytkownik prosi o apkę / APK / plik do pobrania:
1) Pliki projektu NIE są efektem końcowym. Efektem końcowym jest zbudowane app-debug.apk. Każdą odpowiedź zawierającą projekt Android MUSISZ zakończyć zleceniem:
[GH_BUILD: android]
[/GH_BUILD]
Bez tego tagu użytkownik nie dostanie karty buildu ani gotowej apki – same pliki są dla niego bezużyteczne.
2) Kompletny projekt wysyłaj ZAWSZE w jednej odpowiedzi, wszystkie pliki naraz: settings.gradle, build.gradle, app/build.gradle, app/src/main/AndroidManifest.xml, app/src/main/res/values/strings.xml, MainActivity.java oraz pliki assets (w tym assets/call.html – złożony z części, jeśli użytkownik dostarczył go w częściach). NIGDY nie dosyłaj projektu po 1–2 pliki w kolejnych wiadomościach.
3) Nie wysyłaj w kółko tych samych plików. Jeśli poprzednia odpowiedź była niekompletna, wyślij JEDEN kompletny, poprawiony zestaw i zlecenie buildu z pkt 1.
4) Nie zmieniaj architektury w trakcie zadania i nie dokładaj elementów spoza prośby użytkownika (np. lokalnego serwera HTTP). Dla apki z HTML obowiązuje schemat z tej instrukcji: WebView ładujący file:///android_asset/... + onPermissionRequest dla mikrofonu i kamery.
5) Jako „wynik pracy” zakazane są: drzewka katalogów, streszczenia zamiast plików, placeholdery zamiast prawdziwej zawartości oraz obietnice bez zlecenia buildu.
6) Jeśli buildu nie da się zlecić (np. brak konfiguracji GitHub na serwerze), powiedz to użytkownikowi JEDNYM zdaniem wprost – nie maskuj tego kolejnymi porcjami plików projektu.

═════════════════════════════════════
DODATEK AKUMULATOR_V3 – BUDOWANIE NA RATY (OSZCZĘDZANIE TOKENÓW)
═════════════════════════════════════
Ten dodatek jest nadrzędny wobec wymogu „kompletny projekt w jednej odpowiedzi” (również wobec pkt 1–2 dodatku SZKOLENIE_V2) – przy buildach w chmurze obowiązuje nowy tryb:
- Backend zbiera pliki [FILE:] z CAŁEJ rozmowy (zapisuje je jako załączniki Twoich odpowiedzi). Wysyłaj więc projekt małymi partiami: 1–2 pliki na odpowiedź, a gdy zestaw będzie kompletny, wyślij sam tag [GH_BUILD: target] [/GH_BUILD] – może być w osobnej, krótkiej odpowiedzi bez żadnych plików.
- Plików już wysłanych w tej rozmowie NIE powtarzaj (chyba że je poprawiasz – nowa wersja nadpisuje starą). Po każdej partii dopisz jedno zdanie: co już masz wysłane, a czego jeszcze brakuje.
- DUŻYCH plików użytkownika (np. call.html) NIGDY nie przepisuj i nie wysyłaj jako [FILE:] – backend doklei je do projektu z załączników użytkownika co do bajta. Przy apce Android wysyłasz tylko szkielet: settings.gradle, build.gradle, app/build.gradle, app/src/main/AndroidManifest.xml, app/src/main/res/values/strings.xml, MainActivity.java.
- Trzymaj odpowiedzi KRÓTKIE. Mała odpowiedź mieści się w limitach większej liczby kluczy API i rzadziej jest ucinana.

DODATEK AKUMULATOR_V4: treści załączników call.html oraz call-czesc-*.txt NIE dostajesz w rozmowie (zastępuje je notka systemowa) – to celowe oszczędzanie tokenów. Pliki te backend dokleja do buildu sam, co do bajta. Nie proś użytkownika o wklejanie ich treści ani nie zgaduj, co w nich jest – do zbudowania szkieletu apki ich treść nie jest Ci potrzebna.

DODATEK PORZADKI_V5 \u2013 JEDEN PROJEKT = JEDEN CZYSTY ZESTAW:
- Backend przed ka\u017cdym buildem czy\u015bci folder projektu w repo do DOK\u0141ADNIE tego zestawu plik\u00f3w, kt\u00f3ry wejdzie do buildu. Dlatego: jeden build = jeden projekt. Nie mieszaj plik\u00f3w r\u00f3\u017cnych aplikacji ani pakiet\u00f3w w jednym zleceniu.
- Tu\u017c przed tagiem [GH_BUILD] wypisz w jednym zdaniu finalny sk\u0142ad projektu (list\u0119 plik\u00f3w) \u2013 to ostatnia kontrola, \u017ce zestaw jest kompletny i sp\u00f3jny.
- Rozwi\u0105zania, kt\u00f3re ju\u017c dzia\u0142a\u0142y (np. wzorzec MainActivity WebView z onPermissionRequest), wykorzystuj ponownie jako \u015bwie\u017c\u0105 tre\u015b\u0107 dopasowan\u0105 do bie\u017c\u0105cego projektu. Nie do\u0142\u0105czaj plik\u00f3w \u201ena zapas\u201d ani po poprzednich projektach.

DODATEK AUTOBUILD_V6 – NIE ZATRZYMUJ SIĘ:
- Przy buildach NIE czekaj na użytkownika między partiami plików. Backend sam ponawia zapytanie w ramach tego samego zapytania czatu, więc po prostu wysyłaj kolejne partie, aż zestaw będzie kompletny, a na końcu sam tag [GH_BUILD: target] [/GH_BUILD].
- W odpowiedziach z partiami pisz TYLKO pliki [FILE:] – bez zdań wstępu, podsumowań i list „mam/brakuje”. To już niepotrzebne i tylko zużywa tokeny.
- Jeśli backend dopyta mimo kompletnego zestawu, wyślij sam tag buildu – niczego więcej nie trzeba.
"""

# --- POMOCNICZE: MIME i rozszerzenia ---
EXT_MIME = {
    ".py": "text/x-python",
    ".html": "text/html",
    ".htm": "text/html",
    ".js": "application/javascript",
    ".css": "text/css",
    ".json": "application/json",
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".sql": "application/sql",
    ".yaml": "text/yaml",
    ".yml": "text/yaml",
    ".xml": "application/xml",
    ".csv": "text/csv",
    ".sh": "application/x-sh",
    ".ts": "application/typescript",
    ".tsx": "application/typescript",
    ".jsx": "application/javascript",
    ".java": "text/x-java-source",
    ".c": "text/x-c",
    ".cpp": "text/x-c++",
    ".h": "text/x-c",
    ".rs": "text/x-rust",
    ".go": "text/x-go",
    ".php": "application/x-httpd-php",
    ".rb": "application/x-ruby",
    ".swift": "text/x-swift",
    ".kt": "text/x-kotlin",
    ".r": "text/x-r",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}

def guess_mime(filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    return EXT_MIME.get(ext) or mimetypes.guess_type(filename)[0] or "application/octet-stream"

def make_file_response(filename: str, content: bytes | str, mime: str | None = None) -> Dict[str, str]:
    if isinstance(content, str):
        content = content.encode("utf-8")
    return {
        "filename": filename,
        "mime_type": mime or guess_mime(filename),
        "data": base64.b64encode(content).decode("ascii"),
    }

# --- GENERATORY DOCX / PDF ---
def create_docx_from_text(title: str, body: str) -> bytes:
    """Tworzy profesjonalny dokument .docx z prostego tekstu/markdownu."""
    doc = Document()

    # Tytuł
    heading = doc.add_heading(title or "Dokument", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = body.strip().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- ") or line.startswith("* "):
            # lista punktowana
            p = doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(re.sub(r"^\d+\.\s*", "", line), style="List Number")
        else:
            # zwykły akapit – proste **pogrubienie** i *kursywa*
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                    run = p.add_run(part[1:-1])
                    run.italic = True
                else:
                    p.add_run(part)
        i += 1

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _ensure_dejavu_fonts() -> tuple:
    """Pobiera DejaVuSans (regular + bold) do /tmp jeśli brak lokalnych fontów."""
    fonts_dir = os.path.join(os.path.dirname(__file__), "fonts")
    cache_dir = os.path.join("/tmp", "mybro_fonts")
    os.makedirs(cache_dir, exist_ok=True)

    regular_candidates = [
        os.path.join(fonts_dir, "DejaVuSans.ttf"),
        os.path.join(cache_dir, "DejaVuSans.ttf"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ]
    bold_candidates = [
        os.path.join(fonts_dir, "DejaVuSans-Bold.ttf"),
        os.path.join(cache_dir, "DejaVuSans-Bold.ttf"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
    ]

    regular = next((p for p in regular_candidates if os.path.exists(p)), None)
    bold = next((p for p in bold_candidates if os.path.exists(p)), None)

    if not regular:
        urls = [
            "https://github.com/dejavu-fonts/dejavu-fonts/raw/version_2_37/ttf/DejaVuSans.ttf",
            "https://cdn.jsdelivr.net/npm/dejavu-fonts-ttf@2.37.3/ttf/DejaVuSans.ttf",
        ]
        dest = os.path.join(cache_dir, "DejaVuSans.ttf")
        for url in urls:
            try:
                r = requests.get(url, timeout=30)
                r.raise_for_status()
                if len(r.content) > 10000:
                    with open(dest, "wb") as f:
                        f.write(r.content)
                    regular = dest
                    print(f"Downloaded DejaVuSans -> {dest}")
                    break
            except Exception as e:
                print(f"Font download fail ({url}): {e}")

    if not bold:
        urls = [
            "https://github.com/dejavu-fonts/dejavu-fonts/raw/version_2_37/ttf/DejaVuSans-Bold.ttf",
            "https://cdn.jsdelivr.net/npm/dejavu-fonts-ttf@2.37.3/ttf/DejaVuSans-Bold.ttf",
        ]
        dest = os.path.join(cache_dir, "DejaVuSans-Bold.ttf")
        for url in urls:
            try:
                r = requests.get(url, timeout=30)
                r.raise_for_status()
                if len(r.content) > 10000:
                    with open(dest, "wb") as f:
                        f.write(r.content)
                    bold = dest
                    print(f"Downloaded DejaVuSans-Bold -> {dest}")
                    break
            except Exception as e:
                print(f"Font bold download fail ({url}): {e}")

    return regular, bold or regular


def create_pdf_from_text(title: str, body: str) -> bytes:
    """Tworzy PDF z pełną obsługą polskich znaków (DejaVu)."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(left=15, top=15, right=15)
    pdf.add_page()

    regular, bold = _ensure_dejavu_fonts()
    font_loaded = False
    if regular:
        try:
            pdf.add_font("DejaVu", "", regular)
            if bold and bold != regular:
                pdf.add_font("DejaVu", "B", bold)
            else:
                pdf.add_font("DejaVu", "B", regular)
            font_loaded = True
        except Exception as e:
            print(f"Font register error: {e}")

    def set_font(style="", size=11):
        if font_loaded:
            pdf.set_font("DejaVu", style if style in ("", "B") else "", size)
        else:
            pdf.set_font("Helvetica", style if style in ("", "B") else "", size)

    def safe_text(s):
        if not s:
            return ""
        if font_loaded:
            return s
        table = str.maketrans({
            "ą": "a", "ć": "c", "ę": "e", "ł": "l", "ń": "n", "ó": "o", "ś": "s", "ź": "z", "ż": "z",
            "Ą": "A", "Ć": "C", "Ę": "E", "Ł": "L", "Ń": "N", "Ó": "O", "Ś": "S", "Ź": "Z", "Ż": "Z",
        })
        return s.translate(table)

    def write_line(text: str, h: float = 6, align: str = "L"):
        """Bezpieczne multi_cell – zawsze od lewej krawędzi, pełna szerokość."""
        text = safe_text(text or "")
        if not text.strip():
            pdf.ln(h * 0.6)
            return
        pdf.set_x(pdf.l_margin)
        w = getattr(pdf, "epw", None)
        if w is None or w <= 0:
            w = pdf.w - pdf.l_margin - pdf.r_margin
        try:
            pdf.multi_cell(w=w, h=h, text=text, align=align, new_x="LMARGIN", new_y="NEXT")
        except TypeError:
            try:
                pdf.multi_cell(w, h, text, align=align)
            except Exception:
                pdf.multi_cell(w, h, text.encode("ascii", "ignore").decode("ascii"), align=align)
        except Exception:
            try:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(w, h, text.encode("ascii", "ignore").decode("ascii"), align=align)
            except Exception as e2:
                print(f"PDF write_line fallback fail: {e2}")
                pdf.ln(h)

    set_font("B", 16)
    write_line(title or "Dokument", h=10, align="C")
    pdf.ln(4)
    set_font("", 11)

    for line in (body or "").strip().split("\n"):
        line = line.rstrip()
        if not line:
            pdf.ln(4)
            continue
        if line.startswith("# "):
            set_font("B", 14)
            write_line(line[2:].strip(), h=8)
            set_font("", 11)
            pdf.ln(2)
        elif line.startswith("## "):
            set_font("B", 12)
            write_line(line[3:].strip(), h=7)
            set_font("", 11)
            pdf.ln(1)
        elif line.startswith("### "):
            set_font("B", 11)
            write_line(line[4:].strip(), h=6)
            set_font("", 11)
            pdf.ln(1)
        elif line.startswith("- ") or line.startswith("* "):
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", line[2:])
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean)
            write_line("• " + clean.strip(), h=6)
        elif re.match(r"^\d+\.\s", line):
            clean = re.sub(r"^\d+\.\s*", "", line)
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", clean)
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean)
            write_line(clean.strip(), h=6)
        else:
            clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
            clean = re.sub(r"\*([^*]+)\*", r"\1", clean)
            write_line(clean, h=6)

    out = pdf.output()
    if isinstance(out, str):
        out = out.encode("latin-1")
    elif isinstance(out, bytearray):
        out = bytes(out)
    return out



# --- NARZĘDZIA (obrazy, search itd.) ---

def find_attachment_in_history(name, history):
    for item in reversed(history):
        if item.attachments:
            for att in item.attachments:
                if hasattr(att, 'name') and att.name == name: return att
    return None

def open_url_and_get_content(url):
    try:
        headers = {'User-Agent': 'Mozilla/5.0'}
        r = requests.get(url, headers=headers, timeout=20, stream=True)
        r.raise_for_status()
        ct = r.headers.get('Content-Type', '').lower()
        name = url.split('/')[-1].split('?')[0] or "file"
        data = base64.b64encode(r.content).decode()
        if 'image' in ct: return {"type": "image", "filename": f"{name}.png", "data": data}
        elif 'pdf' in ct: return {"type": "pdf", "filename": f"{name}.pdf", "data": data}
        else: raise ValueError("Unsupported file type")
    except Exception as e: raise RuntimeError(f"URL Error: {e}")

def _enhance_image_prompt(prompt: str) -> str:
    prompt = (prompt or "").strip()
    if not prompt:
        return prompt
    low = prompt.lower()
    quality_boost = ", highly detailed, sharp focus, natural lighting"
    if "highly detailed" not in low and "8k" not in low and len(prompt) < 120:
        prompt = prompt.rstrip(",. ") + quality_boost
    return prompt

def _generate_pollinations(prompt: str) -> bytes:
    """Próbuje kilka wariantów URL Pollinations (darmowe parametry najpierw)."""
    enhanced = _enhance_image_prompt(prompt)
    encoded = requests.utils.quote(enhanced)
    # kolejne warianty od najbardziej kompatybilnego / darmowego
    variants = [
        f"https://image.pollinations.ai/prompt/{encoded}?width=1024&height=1024&nologo=true",
        f"https://image.pollinations.ai/prompt/{encoded}?width=1024&height=1024&nologo=true&model=flux",
        f"https://image.pollinations.ai/prompt/{encoded}?width=768&height=768&nologo=true",
        f"https://gen.pollinations.ai/image/{encoded}?width=1024&height=1024&nologo=true",
    ]
    last_err = None
    for url in variants:
        try:
            print(f"ImageGen Pollinations try: {url[:100]}…")
            r = requests.get(url, timeout=90, headers={"User-Agent": "Mozilla/5.0"})
            r.raise_for_status()
            if r.content and len(r.content) >= 1000 and r.headers.get("content-type", "").startswith("image"):
                return r.content
            # czasem zwraca HTML przy limicie
            if r.content and len(r.content) >= 1000 and r.content[:8] != b"<!DOCTYPE":
                return r.content
            last_err = RuntimeError(f"empty/invalid response ({len(r.content or b'')} B)")
        except Exception as e:
            last_err = e
            print(f"ImageGen Pollinations variant fail: {e}")
            continue
    raise RuntimeError(f"Pollinations all variants failed: {last_err}")

def _generate_ai_horde(prompt: str) -> bytes:
    enhanced = _enhance_image_prompt(prompt)
    headers = {
        "apikey": "0000000000",
        "Client-Agent": "MyBro:1.0:eosforus",
        "Content-Type": "application/json"
    }
    # mniejsze wymiary / mniej kroków = mniej odrzuceń na anonimowym kluczu
    payload = {
        "prompt": enhanced,
        "params": {
            "width": 512,
            "height": 512,
            "steps": 20,
            "cfg_scale": 7.0,
            "sampler_name": "k_euler_a",
            "n": 1
        },
        "nsfw": False,
        "censor_nsfw": True,
        "models": ["stable_diffusion"],
        "r2": True
    }
    r = requests.post("https://stablehorde.net/api/v2/generate/async", json=payload, headers=headers, timeout=30)
    r.raise_for_status()
    job_id = r.json().get("id")
    if not job_id:
        raise RuntimeError("AI Horde no job id")
    for _ in range(60):
        time.sleep(3)
        check = requests.get(f"https://stablehorde.net/api/v2/generate/check/{job_id}", headers=headers, timeout=15)
        check.raise_for_status()
        if check.json().get("done"):
            break
        if check.json().get("faulted"):
            raise RuntimeError("AI Horde job faulted")
    else:
        raise RuntimeError("AI Horde timeout")
    result = requests.get(f"https://stablehorde.net/api/v2/generate/status/{job_id}", headers=headers, timeout=30)
    result.raise_for_status()
    gens = result.json().get("generations", [])
    if not gens or not gens[0].get("img"):
        raise RuntimeError("AI Horde no image")
    img_url = gens[0]["img"]
    # R2 URL albo base64
    if img_url.startswith("http"):
        img_resp = requests.get(img_url, timeout=30)
        img_resp.raise_for_status()
        return img_resp.content
    return base64.b64decode(img_url)

def _generate_pollinations_simple(prompt: str) -> bytes:
    """Ostatnia deska ratunku – minimalny endpoint."""
    url = "https://image.pollinations.ai/prompt/" + requests.utils.quote(prompt[:500])
    r = requests.get(url, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    if not r.content or len(r.content) < 500:
        raise RuntimeError("Pollinations simple empty")
    return r.content

def generate_image_from_prompt(prompt: str) -> bytes:
    prompt = (prompt or "").strip()
    if not prompt:
        raise ValueError("Empty prompt")
    if len(prompt) > 1500:
        prompt = prompt[:1500]
    errors = []
    for name, fn in (
        ("Pollinations", _generate_pollinations),
        ("AI Horde", _generate_ai_horde),
        ("Pollinations-simple", _generate_pollinations_simple),
    ):
        try:
            print(f"ImageGen: Trying {name}… prompt={prompt[:80]!r}")
            return fn(prompt)
        except Exception as e:
            errors.append(f"{name}: {e}")
            print(f"ImageGen {name} fail: {e}")
    raise RuntimeError(f"All image engines failed: {'; '.join(errors)}")

def composite_images(fg_bytes, bg_bytes):
    fg = Image.open(BytesIO(fg_bytes)).convert("RGBA")
    bg = Image.open(BytesIO(bg_bytes)).convert("RGBA")
    if fg.width > bg.width or fg.height > bg.height: fg.thumbnail(bg.size, Image.Resampling.LANCZOS)
    x = (bg.width - fg.width) // 2; y = (bg.height - fg.height) // 2
    bg.paste(fg, (x, y), fg)
    buf = BytesIO(); bg.save(buf, format="PNG"); return buf.getvalue()

def edit_image(b64, details):
    img = Image.open(BytesIO(base64.b64decode(b64)))
    op_match = re.search(r'operation="([^"]+)"', details)
    val_match = re.search(r'value="([^"]+)"', details)
    op = (op_match.group(1) if op_match else "show").lower()
    val = val_match.group(1) if val_match else "0"
    if op == "remove_background": 
        return remove(base64.b64decode(b64))
    if op == "show": 
        return base64.b64decode(b64)
    if op == "grayscale": 
        out = img.convert('L')
    elif op == "rotate": 
        out = img.rotate(int(val or 90), expand=True)
    elif op == "blur": 
        out = img.filter(ImageFilter.GaussianBlur(radius=int(val or 2)))
    else: 
        out = img
    buf = BytesIO(); out.save(buf, format="PNG"); return buf.getvalue()

def inpaint_image(b64, prompt):
    if not ZMO_API_KEY: raise ValueError("ZMO Key missing")
    files = {'image': ('i.png', base64.b64decode(b64), 'image/png')}
    r = requests.post(INPAINTING_API_URL, headers={'X-Api-Key': ZMO_API_KEY}, files=files, data={'mask_prompt': prompt}, timeout=90)
    r.raise_for_status(); return r.content

def create_story_pdf(struct, hist):
    pdf = FPDF(); pdf.add_page(); pdf.set_auto_page_break(True, 15)
    try:
        font = os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf")
        pdf.add_font('DejaVu', '', font, uni=True); pdf.set_font('DejaVu', '', 12)
    except: pdf.set_font('Helvetica', '', 12)
    for s in struct:
        if pdf.get_y()>250: pdf.add_page()
        if s.get("text"): pdf.write(8, s.get("text")); pdf.ln(10)
        if s.get("image_prompt"):
            try:
                img = generate_image_from_prompt(s.get("image_prompt"))
                if pdf.get_y()>200: pdf.add_page()
                pdf.image(BytesIO(img), x=15, w=pdf.w-30, type='PNG'); pdf.ln(10)
            except: pass
    out = pdf.output()
    if isinstance(out, str):
        out = out.encode('latin-1')
    elif isinstance(out, bytearray):
        out = bytes(out)
    return out

def search_web(q):
    try:
        with DDGS() as ddgs: return "\n".join([f"{r['title']}: {r['body']}" for r in ddgs.text(q, max_results=4)])
    except Exception as e: return str(e)

def scrape_web(url):
    try:
        r = requests.get(url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=10)
        soup = BeautifulSoup(r.text, 'html.parser')
        for s in soup(["script", "style"]): s.decompose()
        return soup.body.get_text(separator='\n', strip=True)[:8000]
    except Exception as e: return str(e)

def process_att(att):
    """Uniwersalna obsługa DOWOLNYCH załączników: obrazy, PDF, DOCX, kod, tekst, binarki."""
    try:
        raw = base64.b64decode(att.data)
    except Exception:
        return f"[Plik {getattr(att, 'name', 'plik')}: błąd dekodowania base64]"

    mime = (getattr(att, 'mime_type', '') or '').lower()
    name = getattr(att, 'name', '') or 'plik'
    ext = os.path.splitext(name.lower())[1]

    # 1. Obrazy -> obiekt PIL (dla modeli wizyjnych)
    if "image" in mime or ext in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"):
        try:
            return Image.open(BytesIO(raw))
        except Exception as e:
            return f"[Obraz {name}: nie można otworzyć - {e}]"

    # 2. PDF -> wyciągnij tekst (pypdf)
    if mime == "application/pdf" or ext == ".pdf":
        try:
            reader = pypdf.PdfReader(BytesIO(raw))
            texts = []
            for i, page in enumerate(reader.pages[:25]):  # limit 25 stron
                try:
                    t = page.extract_text() or ""
                    if t.strip():
                        texts.append(f"--- Strona {i+1} ---\n{t[:5000]}")
                except Exception:
                    continue
            content = "\n".join(texts)[:18000]
            return f"[Plik PDF: {name}]\n{content or '(nie udało się wyodrębnić tekstu)'}"
        except Exception as e:
            return f"[PDF {name}: błąd odczytu - {e}]"

    # 3. DOCX -> wyciągnij tekst (python-docx)
    if "wordprocessingml" in mime or ext == ".docx":
        try:
            doc = Document(BytesIO(raw))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            # też tabele
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        paras.append(" | ".join(cells))
            content = "\n".join(paras)[:18000]
            return f"[Plik DOCX: {name}]\n{content or '(pusty dokument)'}"
        except Exception as e:
            return f"[DOCX {name}: błąd odczytu - {e}]"

    # 4. Pliki tekstowe / kod -> dekoduj z nazwą i kolorowaniem
    TEXT_EXTS = {
        ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css",
        ".json", ".md", ".markdown", ".txt", ".csv", ".yaml", ".yml",
        ".xml", ".sql", ".sh", ".java", ".c", ".cpp", ".h", ".hpp",
        ".rs", ".go", ".php", ".rb", ".swift", ".kt", ".r", ".toml",
        ".ini", ".env", ".log", ".tex",
    }
    is_text_mime = (
        mime.startswith("text/")
        or mime in ("application/json", "application/xml", "application/javascript",
                    "application/x-python", "application/sql", "application/yaml")
    )
    if ext in TEXT_EXTS or is_text_mime:
        try:
            text = raw.decode("utf-8", errors="ignore")
            truncated = len(text) > 25000
            if truncated:
                text = text[:25000] + f"\n...[ucięto - pełny plik ma {len(raw)} bajtów]"
            lang = ext.lstrip(".") if ext else ""
            return f"[Plik: {name}]\n```{lang}\n{text}\n```"
        except Exception as e:
            return f"[Plik {name}: błąd dekodowania - {e}]"

    # 5. Próba odczytu jako tekst (pliki bez rozszerzenia itd.)
    try:
        sample = raw[:2000]
        if b"\x00" in sample:
            raise ValueError("binary")
        text = raw.decode("utf-8", errors="strict")
        if len(text.strip()) < 5:
            raise ValueError("too short")
        truncated = len(text) > 25000
        if truncated:
            text = text[:25000] + "...[ucięto]"
        return f"[Plik: {name} ({mime or 'nieznany typ'}, {len(raw)} bajtów)]\n{text}"
    except Exception:
        pass

    # 6. Binarka -> metadane + prośba o kontekst
    size_kb = len(raw) / 1024
    size_str = f"{size_kb:.1f} KB" if size_kb < 1024 else f"{size_kb/1024:.1f} MB"
    return (
        f"[Plik binarny: {name} ({mime or 'nieznany typ'}, {size_str}) - "
        f"nie mogę odczytać treści bezpośrednio. "
        f"Opisz, co mam z tym plikiem zrobić.]"
    )

# --- AUTH (JWT, dokładnie tak samo jak w Boxie) ---
# Kody dostępu i sekrety trzymane TYLKO na backendzie, nigdy we frontendzie.
# Na HF Space ustaw w Secrets: JWT_SECRET_KEY (losowy, długi),
# BOX_ACCESS_CODE, BOX_BACKEND_URL.
# SESJA_V8: gdy brak sekretu Space, klucz JWT jest trzymany w datasecie
# (ten sam, co darowane klucze) — dzieki temu restarty Space'a nie uniewazniaja
# logowan. Sekret Space (JWT_SECRET_KEY) nadal ma pierwszenstwo, gdy jest ustawiony.
JWT_SECRET_FILE = "buddy-jwt-secret.json"

def _resolve_jwt_secret():
    env_key = os.environ.get("JWT_SECRET_KEY")
    if env_key:
        return env_key
    if HF_TOKEN:
        try:
            path = hf_hub_download(repo_id=DONATIONS_REPO_ID, filename=JWT_SECRET_FILE,
                                   repo_type="dataset", token=HF_TOKEN, cache_dir=HF_CACHE_DIR)
            with open(path, "r", encoding="utf-8") as f:
                stored = json.load(f).get("secret")
            if stored:
                print("SYSTEM: JWT - wczytano staly klucz z datasetu (SESJA_V8).")
                return stored
        except Exception:
            pass
        generated = secrets.token_urlsafe(32)
        try:
            temp_path = os.path.join(HF_CACHE_DIR, JWT_SECRET_FILE)
            with open(temp_path, "w", encoding="utf-8") as f:
                json.dump({"secret": generated}, f)
            HfApi(token=HF_TOKEN).upload_file(
                path_or_fileobj=temp_path, path_in_repo=JWT_SECRET_FILE,
                repo_id=DONATIONS_REPO_ID, repo_type="dataset",
                commit_message="Add buddy JWT secret")
            print("SYSTEM: JWT - zapisano nowy staly klucz do datasetu (SESJA_V8).")
            return generated
        except Exception as e:
            print(f"OSTRZEZENIE: nie udalo sie zapisac klucza JWT do datasetu ({e}) - "
                  "tokeny uniewaznia sie po restarcie. Ustaw JWT_SECRET_KEY w sekretach Space.")
            return generated
    generated = secrets.token_urlsafe(32)
    print("OSTRZEZENIE: brak JWT_SECRET_KEY i dostepu do datasetu - wygenerowano losowy klucz "
          "(tokeny uniewaznia sie po restarcie). Ustaw JWT_SECRET_KEY w sekretach Space.")
    return generated

JWT_SECRET_KEY = _resolve_jwt_secret()

ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24 * 30  # SESJA_V8: 30 dni (bylo 8h)

# Kody dostępu do agenta (jak w Boxie: kod -> poziom dostępu)
ACCESS_CODES = {
    "2017": "full",
}

BOX_ACCESS_CODE = os.getenv("BOX_ACCESS_CODE", "2358")
BOX_BACKEND_URL = os.getenv("BOX_BACKEND_URL", "https://eosforus-box.hf.space")

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="token", auto_error=False)

class LoginRequest(BaseModel):
    code: str

def create_access_token(data: dict, expires_delta: Union[timedelta, None] = None):
    to_encode = data.copy()
    expire = datetime.utcnow() + (expires_delta or timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, JWT_SECRET_KEY, algorithm=ALGORITHM)

async def get_current_user(request: Request, token_from_header: Optional[str] = Depends(oauth2_scheme)):
    credentials_exception = HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Could not validate credentials")
    token = token_from_header or request.query_params.get("token")
    if not token: raise credentials_exception
    try:
        payload = jwt.decode(token, JWT_SECRET_KEY, algorithms=[ALGORITHM])
        username = payload.get("sub")
        if username != "user":
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    return username

# --- FASTAPI SETUP ---
app = FastAPI()
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

@app.post("/login")
async def login(request: LoginRequest):
    """Logowanie kodem dostępu (jak w Boxie). Zwraca JWT: {"access_token", "token_type"}."""
    code = request.code.strip()
    if code in ACCESS_CODES:
        token = create_access_token(
            data={"sub": "user", "access_level": ACCESS_CODES[code]},
            expires_delta=timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
        )
        return {"access_token": token, "token_type": "bearer"}
    raise HTTPException(status_code=400, detail="Incorrect access code")

class Attachment(BaseModel): name: str; mime_type: str; data: str
class HistoryItem(BaseModel): role: str; content: str; attachments: Optional[List[Attachment]] = None; isHidden: Optional[bool] = False
class ChatRequest(BaseModel): history: List[HistoryItem]; apiKey: Optional[str] = Field(None, alias="apiKey"); request_id: Optional[str] = None
class DonateKeyRequest(BaseModel): apiKey: str
class ExportRequest(BaseModel): history: List[HistoryItem]

@app.post("/donate-key")
def dp(r: DonateKeyRequest, user: str = Depends(get_current_user)):
    k=r.apiKey.strip(); keys=load_donated_keys_safely(); keys.append({"timestamp": datetime.now().isoformat(), "key": k})
    save_donated_keys(keys)
    # dopisz też do puli w pamięci – bez tego klucz działał dla innych dopiero po restarcie backendu
    if k and k not in api_keys: api_keys.append(k); api_keys[:] = _priorytet_kluczy(api_keys)
    return {"status": "ok"}

@app.post("/export-pdf")
def ep(r: ExportRequest, user: str = Depends(get_current_user)):
    pdf = FPDF(); pdf.add_page(); pdf.set_auto_page_break(True, 15)
    try:
        font = os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf")
        pdf.add_font('DejaVu', '', font, uni=True); pdf.set_font('DejaVu', '', 10)
    except: pdf.set_font('Helvetica', '', 10)
    for m in r.history:
        if m.isHidden: continue
        pdf.multi_cell(0, 5, f"{'Ty' if m.role=='user' else 'AI'}: {m.content.replace('[IMG]','')}"); pdf.ln(2)
    out = pdf.output()
    if isinstance(out, str):
        out = out.encode('latin-1')
    elif isinstance(out, bytearray):
        out = bytes(out)
    return StreamingResponse(BytesIO(out), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=chat.pdf"})


# --- BOX PROXY: kod dostępu do Boxa zna tylko backend, frontend woła /box/* ---
_box_token: Optional[str] = None
_box_token_exp: float = 0.0

def _box_login() -> str:
    global _box_token, _box_token_exp
    if _box_token and time.time() < _box_token_exp:
        return _box_token
    if not BOX_ACCESS_CODE:
        raise HTTPException(status_code=503, detail="Box nie jest skonfigurowany na backendzie.")
    try:
        r = requests.post(f"{BOX_BACKEND_URL}/login", json={"code": BOX_ACCESS_CODE}, timeout=15)
        r.raise_for_status()
        data = r.json()
        _box_token = data.get("access_token") or data.get("token")
        if not _box_token:
            raise RuntimeError("Box login: brak tokenu w odpowiedzi")
        _box_token_exp = time.time() + 3000
        return _box_token
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Box login error: {e}")

def _box_headers() -> dict:
    return {"Authorization": f"Bearer {_box_login()}"}

@app.get("/box/files")
def box_files(user: str = Depends(get_current_user)):
    try:
        r = requests.get(f"{BOX_BACKEND_URL}/files", headers=_box_headers(), timeout=20)
        r.raise_for_status()
        return r.json()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Box files error: {e}")

@app.post("/box/upload")
async def box_upload(
    file: UploadFile = File(...),
    path: str = Form("chats"),
    user: str = Depends(get_current_user),
):
    try:
        data = await file.read()
        files = {"file": (file.filename, data, file.content_type or "application/octet-stream")}
        r = requests.post(f"{BOX_BACKEND_URL}/upload", headers=_box_headers(),
                           files=files, data={"path": path}, timeout=60)
        r.raise_for_status()
        try:
            return r.json()
        except Exception:
            return {"status": "ok"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Box upload error: {e}")

@app.get("/box/download")
def box_download(path: str, user: str = Depends(get_current_user)):
    try:
        enc = requests.utils.quote(path, safe="")
        bt = _box_login()
        r = requests.get(f"{BOX_BACKEND_URL}/download/{enc}?token={bt}&preview=true", timeout=30)
        r.raise_for_status()
        return r.json()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Box download error: {e}")

class BoxDeleteRequest(BaseModel):
    path: str

@app.post("/box/delete")
def box_delete(r: BoxDeleteRequest, user: str = Depends(get_current_user)):
    try:
        resp = requests.post(f"{BOX_BACKEND_URL}/delete", headers=_box_headers(),
                             json={"path": r.path}, timeout=20)
        resp.raise_for_status()
        try:
            return resp.json()
        except Exception:
            return {"status": "ok"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Box delete error: {e}")


# --- GŁOS (TTS) ---
class TTSRequest(BaseModel):
    text: str
    voice: Optional[str] = "pl-PL-MarekNeural"

@app.post("/tts")
async def tts(r: TTSRequest, user: str = Depends(get_current_user)):
    try:
        communicate = edge_tts.Communicate(r.text, r.voice)
        audio_bytes = b""
        async for chunk in communicate.stream():
            if chunk["type"] == "audio":
                audio_bytes += chunk["data"]
        return StreamingResponse(
            BytesIO(audio_bytes),
            media_type="audio/mpeg",
            headers={"Content-Disposition": "inline; filename=speech.mp3"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"TTS error: {str(e)}")


# --- GŁOS (STT) ---
def _transcribe_with_key(key: str, audio_path: str, language: Optional[str] = None) -> str:
    lang = language if language in ("pl", "en", "de", "fr", "es", "it", "pt", "ru", "uk", "cs", "sk") else None
    if key.startswith("gsk_"):
        client = groq.Groq(api_key=key)
        with open(audio_path, "rb") as f:
            kwargs = {
                "file": (os.path.basename(audio_path), f),
                "model": "whisper-large-v3",
                "response_format": "text",
            }
            if lang:
                kwargs["language"] = lang
            result = client.audio.transcriptions.create(**kwargs)
        return result if isinstance(result, str) else getattr(result, "text", str(result))
    if key.startswith("sk-") and not key.startswith("sk-ant-"):
        client = openai.OpenAI(api_key=key)
        with open(audio_path, "rb") as f:
            kwargs = {
                "file": f,
                "model": "whisper-1",
                "response_format": "text",
            }
            if lang:
                kwargs["language"] = lang
            result = client.audio.transcriptions.create(**kwargs)
        return result if isinstance(result, str) else getattr(result, "text", str(result))
    raise ValueError("Key type does not support STT (need Groq gsk_ or OpenAI sk-)")


@app.post("/stt")
async def stt(
    file: UploadFile = File(...),
    language: Optional[str] = Form(None),
    apiKey: Optional[str] = Form(None),
    user: str = Depends(get_current_user),
):
    global current_key_index
    raw = await file.read()
    if not raw or len(raw) < 100:
        raise HTTPException(status_code=400, detail="Pusty lub zbyt krótki plik audio.")
    orig_name = file.filename or "audio.webm"
    suffix = "." + orig_name.rsplit(".", 1)[-1].lower() if "." in orig_name else ".webm"
    if suffix not in (".webm", ".mp4", ".m4a", ".wav", ".ogg", ".mp3", ".mpeg", ".mpga", ".oga", ".flac"):
        suffix = ".webm"
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(raw)
            tmp_path = tmp.name
        keys = ([apiKey] if apiKey else []) + api_keys[current_key_index:] + api_keys[:current_key_index]
        keys = [k for k in keys if k and (k.startswith("gsk_") or (k.startswith("sk-") and not k.startswith("sk-ant-")))]
        if not keys:
            raise HTTPException(
                status_code=503,
                detail="Brak kluczy API z obsługą STT (Groq lub OpenAI). Podaj własny klucz gsk_/sk-."
            )
        last_err = None
        for idx, k in enumerate(keys):
            if idx > 8:
                break
            try:
                text = _transcribe_with_key(k, tmp_path, language)
                if k in api_keys:
                    current_key_index = api_keys.index(k)
                return {"text": (text or "").strip()}
            except Exception as e:
                print(f"STT key fail ({k[:8]}...): {e}")
                last_err = e
                continue
        raise HTTPException(status_code=502, detail=f"STT failed: {last_err}")
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


# --- LLM CALL ---
def call_llm(key, msgs, sys_prompt):
    print(f"-> Call: {key[:8]}...")
    
    def fmt(prov):
        out = []
        for m in msgs:
            if m.get('isHidden'): continue
            role = 'assistant' if m['role']=='model' else 'user'
            content = " ".join(str(p) for p in m['parts'] if isinstance(p, str))
            if content: out.append({'role': role, 'content': content})
        if prov=='anthropic' and out and out[0]['role']=='assistant': del out[0]
        return out

    try:
        if key.startswith("sk-ant-"):
            client = anthropic.Anthropic(api_key=key)
            return client.messages.create(
                model="claude-sonnet-5-5",
                max_tokens=8000,
                system=sys_prompt,
                messages=fmt('anthropic')
            ).content[0].text

        elif key.startswith("gsk_"):
            client = groq.Groq(api_key=key)
            models = [
                "openai/gpt-oss-120b",
                "openai/gpt-oss-20b",
                "qwen/qwen3.8-27b",
                "llama-3.3-70b-versatile"
            ]
            payload = fmt('groq')
            payload.insert(0, {"role": "system", "content": sys_prompt})
            for m in models:
                try:
                    print(f"  Trying Groq model: {m}")
                    return client.chat.completions.create(
                        model=m, 
                        messages=payload, 
                        max_tokens=8000
                    ).choices[0].message.content
                except Exception as e:
                    print(f"  Groq {m}: {e}")
                    continue
            raise RuntimeError("All Groq models failed")

        elif key.startswith("LLM_"):
            # Meta Model API (modele Muse) – endpoint zgodny z OpenAI
            client = openai.OpenAI(api_key=key, base_url="https://api.meta.ai/v1")
            models_meta = [
                "muse-spark-1.3",
                "muse-spark-1.1",
            ]
            payload = fmt('meta')
            payload.insert(0, {"role": "system", "content": sys_prompt})
            for m in models_meta:
                try:
                    print(f"  Trying Meta model: {m}")
                    return client.chat.completions.create(
                        model=m,
                        messages=payload,
                        max_tokens=8000
                    ).choices[0].message.content
                except Exception as e:
                    print(f"  Meta {m}: {e}")
                    continue
            raise RuntimeError("All Meta models failed")

        elif key.startswith("sk-or-"):
            # OpenRouter - endpoint zgodny z OpenAI, klucze sk-or-v1-...
            # MUSI być przed ogólną gałęzią "sk-", inaczej klucz trafi do OpenAI (401)
            client = openai.OpenAI(api_key=key, base_url="https://openrouter.ai/api/v1")
            models_or = [
                "openrouter/auto",
                "openai/gpt-4o-mini",
                "meta-llama/llama-3.3-70b-instruct",
            ]
            payload = fmt('openai')
            payload.insert(0, {"role": "system", "content": sys_prompt})
            for m in models_or:
                try:
                    print(f"  Trying OpenRouter model: {m}")
                    return client.chat.completions.create(
                        model=m,
                        messages=payload,
                        max_tokens=8000
                    ).choices[0].message.content
                except Exception as e:
                    print(f"  OpenRouter {m}: {e}")
                    continue
            raise RuntimeError("All OpenRouter models failed")

        elif key.startswith("sk-"):
            client = openai.OpenAI(api_key=key)
            payload = fmt('openai')
            payload.insert(0, {"role": "system", "content": sys_prompt})
            return client.chat.completions.create(
                model="gpt-4o",
                messages=payload,
                max_tokens=8000
            ).choices[0].message.content

        elif key.startswith("AIza"):
            client = genai.Client(api_key=key)
            # PRIORYTET_V7: usuniete gemini-2.5-flash i gemini-2.5-pro — Google
            # zwraca dla nich 404 (niedostepne dla nowych uzytkownikow), wiec tylko
            # marnowaly czas przy kazdym wywolaniu klucza Gemini.
            models_gemini = [
                "gemini-3.8-flash",
                "gemini-3.5-flash",
                "gemini-3.1-pro-preview"
            ]
            contents = []
            for m in msgs:
                if m.get('isHidden'): continue
                role = "user" if m['role'] == 'user' else "model"
                parts = []
                for p in m.get('parts', []):
                    if isinstance(p, str) and p.strip():
                        parts.append(types.Part.from_text(text=p))
                    elif isinstance(p, Image.Image):
                        # google-genai wymaga Part z danymi bajtowymi, nie surowego PIL
                        try:
                            buf = BytesIO()
                            img = p
                            if img.mode not in ("RGB", "RGBA"):
                                img = img.convert("RGB")
                            img.save(buf, format="PNG")
                            parts.append(types.Part.from_bytes(data=buf.getvalue(), mime_type="image/png"))
                        except Exception as img_err:
                            print(f"  Gemini image part skip: {img_err}")
                            parts.append(types.Part.from_text(text="[obraz – nie udało się dołączyć]"))
                if parts:
                    contents.append(types.Content(role=role, parts=parts))
            for gm in models_gemini:
                try:
                    print(f"  Trying Gemini model: {gm}")
                    response = client.models.generate_content(
                        model=gm,
                        contents=contents,
                        config=types.GenerateContentConfig(
                            system_instruction=sys_prompt,
                            max_output_tokens=8000,
                            temperature=0.7
                        )
                    )
                    return response.text
                except Exception as e:
                    print(f"  Gemini {gm}: {e}")
                    continue
            raise RuntimeError("All Gemini models failed")
        else:
            # MISTRAL_V1: klucze Mistral nie maja rozpoznawalnego prefiksu (losowe znaki),
            # wiec kazdy klucz spoza znanych formatow probujemy przez API Mistral
            # (endpoint zgodny z OpenAI). Jesli to nie jest wazny klucz Mistral,
            # API zwroci blad i rotacja kluczy przejdzie dalej - jak przy martwym kluczu.
            client = openai.OpenAI(api_key=key, base_url="https://api.mistral.ai/v1")
            models_mistral = [
                "mistral-large-latest",
                "mistral-medium-latest",
                "mistral-small-latest",
            ]
            payload = fmt('openai')
            payload.insert(0, {"role": "system", "content": sys_prompt})
            for m in models_mistral:
                try:
                    print(f"  Trying Mistral model: {m}")
                    return client.chat.completions.create(
                        model=m,
                        messages=payload,
                        max_tokens=8000
                    ).choices[0].message.content
                except Exception as e:
                    print(f"  Mistral {m}: {e}")
                    continue
            raise RuntimeError("All Mistral models failed")
    except Exception as e:
        raise RuntimeError(f"Provider error: {e}")


# --- BUILDY W CHMURZE: GitHub Actions (konfiguracja w sekretach Space) ---
GH_TOKEN = os.getenv("GITHUB_TOKEN", "")
GH_BUILD_REPO = os.getenv("GH_BUILD_REPO", "")  # np. "Eosforus/agent-builds"
GH_BUILD_WORKFLOW = os.getenv("GH_BUILD_WORKFLOW", "build.yml")
GH_BUILD_REF = os.getenv("GH_BUILD_REF", "main")
GH_BUILD_WAIT_SECONDS = int(os.getenv("GH_BUILD_WAIT_SECONDS", "240"))
GH_TARGETS = {"nsis", "dotnet", "cpp", "android"}

# Treść workflow wgrywanego automatycznie do repo buildowego
# (funkcja _gh_ensure_workflow) – zgodna z plikiem github_actions/build.yml
GH_WORKFLOW_YML = '''name: Agent Builds

# Workflow wywoływany zdalnie przez backend agenta (workflow_dispatch).
# Pliki projektu agent wgrywa do folderu project/ przed wywołaniem.
on:
  workflow_dispatch:
    inputs:
      target:
        description: "Co budujemy"
        required: true
        default: "nsis"
        type: choice
        options:
          - nsis
          - dotnet
          - cpp
          - android

permissions:
  contents: read

jobs:
  build:
    # C# wymaga prawdziwego Windowsa; NSIS i MinGW budujemy taniej na Ubuntu
    runs-on: ${{ inputs.target == 'dotnet' && 'windows-latest' || 'ubuntu-latest' }}
    steps:
      - uses: actions/checkout@v4

      - name: Instalator NSIS
        if: inputs.target == 'nsis'
        run: |
          sudo apt-get update
          sudo apt-get install -y nsis
          cd project
          makensis installer.nsi

      - name: Program C++ (MinGW -> Windows .exe)
        if: inputs.target == 'cpp'
        run: |
          sudo apt-get update
          sudo apt-get install -y g++-mingw-w64-x86-64
          cd project
          x86_64-w64-mingw32-g++ -O2 -static -o app.exe *.cpp

      - name: Program .NET (Windows .exe)
        if: inputs.target == 'dotnet'
        shell: pwsh
        run: |
          dotnet publish project -c Release -r win-x64 --self-contained true -p:PublishSingleFile=true -o project/out

      # Maszyny ubuntu-latest mają Android SDK i JDK 17 fabrycznie;
      # dokładamy tylko Gradle (projekt nie zawiera wrappera binarnego)
      - name: Gradle (tylko Android)
        if: inputs.target == 'android'
        uses: gradle/actions/setup-gradle@v4

      - name: Android APK (WebView, debug)
        if: inputs.target == 'android'
        run: |
          cd project
          gradle assembleDebug --no-daemon

      - name: Artefakt wynikowy
        uses: actions/upload-artifact@v4
        with:
          name: build-output
          path: |
            project/*.exe
            project/out/*.exe
            project/app/build/outputs/apk/debug/*.apk
          if-no-files-found: error
'''


def _gh_api(method, path, **kwargs):
    url = f"https://api.github.com/repos/{GH_BUILD_REPO}{path}"
    headers = {
        "Authorization": f"Bearer {GH_TOKEN}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }
    return requests.request(method, url, headers=headers, timeout=60, **kwargs)

def _gh_safe_relpath(fname: str) -> str:
    """Bezpieczna ścieżka względna pliku projektu (obsługuje podfoldery,
    odrzuca wyjście poza folder project/)."""
    p = fname.replace("\\", "/").lstrip("/")
    parts = [seg for seg in p.split("/") if seg not in ("", ".")]
    if not parts or any(seg == ".." for seg in parts):
        raise RuntimeError(f"Niedozwolona ścieżka pliku projektu: {fname}")
    return "/".join(parts)

def _gh_list_project_files():
    """PORZADKI_V5: wszystkie pliki pod project/ w repo buildowym,
    rekurencyjnie (Git Trees API). Zwraca {sciezka_wzgledna: sha} albo None,
    gdy nie da się pobrać drzewa (wtedy działa stare, płytkie sprzątanie)."""
    r = _gh_api("GET", f"/git/trees/{GH_BUILD_REF}?recursive=1")
    if r.status_code != 200:
        return None
    out = {}
    for it in r.json().get("tree", []):
        p = it.get("path", "")
        if it.get("type") == "blob" and p.startswith("project/"):
            out[p[len("project/"):]] = it.get("sha")
    return out

def _gh_upload_project(sources: dict):
    """Wgrywa pliki projektu do folderu project/ w repo buildowym (Contents API),
    usuwając pliki po poprzednich buildach. Nazwy plików mogą zawierać
    podfoldery (np. app/src/main/AndroidManifest.xml w projektach Androida)."""
    from urllib.parse import quote
    rel_sources = {_gh_safe_relpath(fn): data for fn, data in sources.items()}
    # PORZADKI_V5: pełne sprzątanie. Wcześniej czyszczony był tylko wierzch
    # folderu project/ (pliki z tej samej listy), więc wszystko w podfolderach
    # – stare assets (np. spooky_bat.html) i klasy poprzednich projektów –
    # zostawało w repo na zawsze i lądowało w kolejnych buildach. Teraz
    # usuwamy KAŻDY plik spoza bieżącego zestawu, rekurencyjnie.
    existing_all = _gh_list_project_files()
    if existing_all is not None:
        for rel_old, sha_old in existing_all.items():
            if rel_old not in rel_sources:
                _gh_api("DELETE", f"/contents/project/{quote(rel_old, safe='/')}",
                        json={"message": "agent: cleanup", "sha": sha_old, "branch": GH_BUILD_REF})
    else:
        # awaryjnie: stare, płytkie sprzątanie samego wierzchu folderu
        top_level = {rel.split("/")[0] for rel in rel_sources}
        r = _gh_api("GET", "/contents/project")
        if r.status_code == 200:
            for entry in r.json():
                if entry.get("type") == "file" and entry["name"] not in top_level:
                    _gh_api("DELETE", f"/contents/project/{quote(entry['name'])}",
                            json={"message": "agent: cleanup", "sha": entry["sha"], "branch": GH_BUILD_REF})
    for rel, data in rel_sources.items():
        url_path = quote(rel, safe="/")
        body = {
            "message": f"agent: upload {rel}",
            "content": base64.b64encode(data).decode("ascii"),
            "branch": GH_BUILD_REF,
        }
        existing = _gh_api("GET", f"/contents/project/{url_path}")
        if existing.status_code == 200:
            body["sha"] = existing.json()["sha"]
        resp = _gh_api("PUT", f"/contents/project/{url_path}", json=body)
        if resp.status_code not in (200, 201):
            raise RuntimeError(f"GitHub upload {rel}: HTTP {resp.status_code} {resp.text[:200]}")

def _gh_check_workflow():
    """Sprawdza, czy GitHub Actions widzi workflow build.yml w repo.
    Zamienia goły 404 przy dispatchu na konkretną przyczynę."""
    want = f".github/workflows/{GH_BUILD_WORKFLOW}"
    r = _gh_api("GET", "/actions/workflows")
    if r.status_code != 200:
        raise RuntimeError(
            f"Nie udało się pobrać listy workflowów z GitHuba (HTTP {r.status_code}). "
            "Sprawdź, czy token ma uprawnienie Actions: Read and write i obejmuje to repo.")
    paths = [w.get("path", "") for w in r.json().get("workflows", [])]
    if want not in paths:
        seen = ", ".join(p for p in paths if p) or "brak"
        raise RuntimeError(
            f"GitHub nie widzi workflow '{GH_BUILD_WORKFLOW}' (zarejestrowane workflowy: {seen}). "
            f"Plik musi leżeć w ścieżce {want} na gałęzi {GH_BUILD_REF}. Jeśli leży w głównym "
            "katalogu repo, przenieś go: usuń stamtąd i utwórz nowy, wpisując w nazwie pliku "
            f"pełną ścieżkę {want}.")

def _gh_dispatch(target: str) -> float:
    started = time.time()
    resp = _gh_api("POST", f"/actions/workflows/{GH_BUILD_WORKFLOW}/dispatches",
                   json={"ref": GH_BUILD_REF, "inputs": {"target": target}})
    if resp.status_code != 204:
        raise RuntimeError(f"GitHub dispatch: HTTP {resp.status_code} {resp.text[:200]}")
    return started

def _gh_find_run(started: float) -> int:
    """Dispatch nie zwraca ID, więc szukamy najnowszego runa po czasie startu."""
    for _ in range(12):
        time.sleep(5)
        r = _gh_api("GET", f"/actions/workflows/{GH_BUILD_WORKFLOW}/runs",
                    params={"event": "workflow_dispatch", "per_page": "5"})
        if r.status_code == 200:
            for run in r.json().get("workflow_runs", []):
                created = run.get("created_at", "")
                if created:
                    dt = datetime.strptime(created, "%Y-%m-%dT%H:%M:%SZ")
                    ts = (dt - datetime(1970, 1, 1)).total_seconds()
                    if ts >= started - 10:
                        return run["id"]
    raise RuntimeError("Nie znaleziono uruchomionego builda na GitHubie (sprawdź zakładkę Actions w repo).")

def _gh_run_status(run_id: int) -> dict:
    r = _gh_api("GET", f"/actions/runs/{run_id}")
    if r.status_code != 200:
        raise RuntimeError(f"GitHub status: HTTP {r.status_code}")
    return r.json()

def _gh_download_outputs(run_id: int) -> dict:
    import zipfile
    r = _gh_api("GET", f"/actions/runs/{run_id}/artifacts")
    if r.status_code != 200:
        raise RuntimeError(f"GitHub artifacts: HTTP {r.status_code}")
    arts = [a for a in r.json().get("artifacts", []) if not a.get("expired")]
    if not arts:
        raise RuntimeError("Build nie zostawił artefaktów (workflow musi kończyć się krokiem upload-artifact).")
    z = _gh_api("GET", f"/actions/artifacts/{arts[0]['id']}/zip")
    if z.status_code != 200:
        raise RuntimeError(f"GitHub artifact download: HTTP {z.status_code}")
    out = {}
    with zipfile.ZipFile(BytesIO(z.content)) as zf:
        for name in zf.namelist():
            if not name.endswith("/"):
                out[os.path.basename(name)] = zf.read(name)
    if not out:
        raise RuntimeError("Artefakt z builda jest pusty.")
    return out

def _gh_validate(target: str, sources: dict) -> str:
    """Wspólna walidacja zlecenia buildu; zwraca znormalizowany target."""
    if not GH_TOKEN or not GH_BUILD_REPO:
        raise RuntimeError("Brak konfiguracji GitHub: ustaw sekrety Space GITHUB_TOKEN i GH_BUILD_REPO.")
    target = (target or "nsis").strip().lower()
    if target not in GH_TARGETS:
        raise RuntimeError(f"Nieznany target GH_BUILD: {target} (dozwolone: {sorted(GH_TARGETS)})")
    if not sources:
        raise RuntimeError("Brak plików projektu – agent musi dołączyć je tagami [FILE:] w tej rozmowie.")
    return target

def _gh_check_access() -> None:
    """Wstępna diagnostyka repo buildowego – zamienia gołe 404 z API
    na czytelne komunikaty (zła nazwa repo / token bez dostępu / puste repo)."""
    r = _gh_api("GET", "")
    if r.status_code == 404:
        raise RuntimeError(
            f"GitHub nie widzi repo '{GH_BUILD_REPO}' dla tego tokenu (HTTP 404). "
            "Sprawdź: (1) sekret GH_BUILD_REPO ma format login/agent-builds "
            "(bez https:// i bez .git), (2) token obejmuje to repo "
            "(Repository access przy tokenie), (3) sekret GITHUB_TOKEN zawiera "
            "aktualny token.")
    if r.status_code != 200:
        raise RuntimeError(f"GitHub: sprawdzenie repo zwróciło HTTP {r.status_code}.")
    default_branch = r.json().get("default_branch")
    if not default_branch:
        raise RuntimeError(
            f"Repo '{GH_BUILD_REPO}' jest puste – nie ma jeszcze żadnej gałęzi. "
            "Dodaj w nim na GitHubie pierwszy plik (np. README.md przyciskiem "
            "Add file → Create new file), co utworzy gałąź main, i spróbuj ponownie.")
    if GH_BUILD_REF != default_branch:
        b = _gh_api("GET", f"/branches/{GH_BUILD_REF}")
        if b.status_code != 200:
            raise RuntimeError(
                f"W repo '{GH_BUILD_REPO}' nie ma gałęzi '{GH_BUILD_REF}' "
                f"(domyślna gałąź to '{default_branch}'). Ustaw na Space zmienną "
                f"GH_BUILD_REF={default_branch} albo zmień nazwę domyślnej gałęzi w repo.")

def _gh_ensure_workflow():
    """Dba o to, by w repo buildowym istniał plik .github/workflows/build.yml.
    Gdy go nie ma, wgrywa wbudowaną wersję – wymaga to uprawnienia tokenu
    Workflows: Read and write (GitHub bez niego odrzuca zapis workflow, 403)."""
    path = f".github/workflows/{GH_BUILD_WORKFLOW}"
    if _gh_api("GET", f"/contents/{path}").status_code == 200:
        return
    body = {
        "message": "agent: dodaj workflow build.yml",
        "content": base64.b64encode(GH_WORKFLOW_YML.encode("utf-8")).decode("ascii"),
        "branch": GH_BUILD_REF,
    }
    resp = _gh_api("PUT", f"/contents/{path}", json=body)
    if resp.status_code not in (200, 201):
        raise RuntimeError(
            f"Nie udało się utworzyć {path} w repo (HTTP {resp.status_code}). "
            "Najpewniej token nie ma uprawnienia Workflows: Read and write – "
            "dopisz je w ustawieniach tokenu na GitHubie albo utwórz ten plik ręcznie.")

def github_build_start(target: str, sources: dict) -> tuple[int, str]:
    """Wgrywa projekt i wywołuje workflow. Zwraca (run_id, run_url)."""
    target = _gh_validate(target, sources)
    _gh_check_access()
    _gh_ensure_workflow()
    _gh_upload_project(sources)
    _gh_check_workflow()
    started = _gh_dispatch(target)
    run_id = _gh_find_run(started)
    return run_id, f"https://github.com/{GH_BUILD_REPO}/actions/runs/{run_id}"

def github_build_collect(run_id: int, max_seconds: int) -> dict | None:
    """Czeka na zakończenie runa (polling co 10 s). Zwraca pliki wynikowe,
    None przy przekroczeniu max_seconds; RuntimeError przy niepowodzeniu."""
    run_url = f"https://github.com/{GH_BUILD_REPO}/actions/runs/{run_id}"
    deadline = time.time() + max_seconds
    while True:
        run = _gh_run_status(run_id)
        if run.get("status") == "completed":
            if run.get("conclusion") == "success":
                return _gh_download_outputs(run_id)
            raise RuntimeError(f"Build #{run_id} nieudany (wynik: {run.get('conclusion')}). Logi: {run_url}")
        if time.time() > deadline:
            return None
        time.sleep(10)

def github_build(target: str, sources: dict) -> tuple[dict, str]:
    """Wariant synchroniczny (start + oczekiwanie do GH_BUILD_WAIT_SECONDS)."""
    run_id, run_url = github_build_start(target, sources)
    outputs = github_build_collect(run_id, GH_BUILD_WAIT_SECONDS)
    if outputs is None:
        return {}, (f"Build #{run_id} nadal trwa (limit oczekiwania {GH_BUILD_WAIT_SECONDS} s). "
                    f"Przebieg: {run_url} . Gdy użytkownik zapyta o ten build, dokończysz go "
                    f"tagiem [GH_BUILD_STATUS: {run_id}].")
    return outputs, f"Build #{run_id} zakończony sukcesem. {run_url}"

def github_build_status(run_id: int) -> tuple[dict, str]:
    """Sprawdza wcześniejszy build; zwraca jego pliki, gdy już się zakończył."""
    if not GH_TOKEN or not GH_BUILD_REPO:
        raise RuntimeError("Brak konfiguracji GitHub: ustaw sekrety Space GITHUB_TOKEN i GH_BUILD_REPO.")
    run = _gh_run_status(run_id)
    run_url = f"https://github.com/{GH_BUILD_REPO}/actions/runs/{run_id}"
    if run.get("status") != "completed":
        return {}, f"Build #{run_id} nadal trwa (status: {run.get('status')}). {run_url}"
    if run.get("conclusion") != "success":
        raise RuntimeError(f"Build #{run_id} nieudany (wynik: {run.get('conclusion')}). Logi: {run_url}")
    return _gh_download_outputs(run_id), f"Build #{run_id} zakończony sukcesem. {run_url}"


# --- Asynchroniczne buildy dla karty w czacie (frontend odpytuje o status) ---
import threading
import uuid

CLOUD_BUILDS: dict = {}
GH_BUILD_MAX_SECONDS = int(os.getenv("GH_BUILD_MAX_SECONDS", "1800"))

def _cloud_build_worker(job_id: str, target: str, sources: dict):
    job = CLOUD_BUILDS[job_id]
    try:
        job["stage"] = "Wgrywam pliki projektu na GitHuba…"
        run_id, run_url = github_build_start(target, sources)
        job["run_id"] = run_id
        job["run_url"] = run_url
        job["stage"] = "Build w toku na GitHubie…"
        outputs = github_build_collect(run_id, GH_BUILD_MAX_SECONDS)
        if outputs is None:
            raise RuntimeError(f"Build przekroczył limit {GH_BUILD_MAX_SECONDS} s. Przebieg: {run_url}")
        job["files"] = [make_file_response(n, d, "application/octet-stream") for n, d in outputs.items()]
        job["status"] = "done"
        job["stage"] = "Gotowe"
    except Exception as e:
        job["status"] = "failed"
        job["error"] = str(e)
        job["stage"] = "Build nieudany"

def create_cloud_build(target: str, sources: dict) -> str:
    """Waliduje zlecenie i startuje build w wątku w tle. Zwraca job_id."""
    _gh_validate(target, sources)
    job_id = uuid.uuid4().hex[:12]
    CLOUD_BUILDS[job_id] = {
        "status": "running", "stage": "Startuję build…", "target": target.strip().lower(),
        "run_id": None, "run_url": None, "files": [], "error": None, "created": time.time(),
    }
    threading.Thread(target=_cloud_build_worker, args=(job_id, target, sources), daemon=True).start()
    return job_id

def cloud_build_snapshot(job_id: str) -> dict | None:
    job = CLOUD_BUILDS.get(job_id)
    if not job:
        return None
    snap = {"status": job["status"], "stage": job["stage"], "target": job["target"],
            "run_url": job["run_url"], "error": job["error"]}
    if job["status"] == "done":
        snap["files"] = job["files"]
    return snap

class CloudBuildRequest(BaseModel):
    target: str
    sources: List[Dict[str, str]]  # [{filename, data(base64)}]

@app.post("/cloud-build/start")
def cloud_build_start_ep(r: CloudBuildRequest, user: str = Depends(get_current_user)):
    sources = {}
    for s in r.sources:
        fn = s.get("filename", "")
        try:
            sources[fn] = base64.b64decode(s.get("data", ""))
        except Exception:
            raise HTTPException(status_code=400, detail=f"Nieprawidłowe dane pliku: {fn}")
    try:
        job_id = create_cloud_build(r.target, sources)
    except RuntimeError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"build_id": job_id, "status": "running"}

@app.get("/cloud-build/{job_id}")
def cloud_build_status_ep(job_id: str, user: str = Depends(get_current_user)):
    snap = cloud_build_snapshot(job_id)
    if snap is None:
        raise HTTPException(status_code=404, detail="Build nie istnieje (mógł wygasnąć przy restarcie serwera).")
    return snap


# --- Sterowanie aktywnym zadaniem czatu z chmurki oczekiwania ---
# Prosta odpowiedź = frontend pokazuje tylko trzy kropki. Gdy odpowiedź trwa
# długo (duży projekt: backend dokleja kolejne części w pętli), chmurka
# przechodzi w tryb rozbudowany: czas, szacunkowy procent oraz przyciski
# pauza / wznów / stop. Przyciski wołają te endpointy, a pętla w /chat
# sprawdza flagi między częściami – pauza i stop są więc prawdziwe.
CHAT_JOBS: dict = {}

def _chat_job_register(rid: str):
    if not rid:
        return
    now = time.time()
    for k in [k for k, v in CHAT_JOBS.items() if now - v.get("started", now) > 1800]:
        CHAT_JOBS.pop(k, None)  # sprzątanie po starych zadaniach
    CHAT_JOBS[rid] = {"paused": False, "stop": False, "stage": "Pracuję nad odpowiedzią…",
                      "parts": 0, "started": now}

def _chat_job_set_stage(rid: str | None, stage: str, parts: int | None = None):
    job = CHAT_JOBS.get(rid or "")
    if job is None:
        return
    job["stage"] = stage
    if parts is not None:
        job["parts"] = parts

def _chat_job_gate(rid: str | None) -> bool:
    """Czeka, gdy użytkownik wstrzymał zadanie z chmurki. Zwraca False, gdy
    zażądano stopu – pętla ma wtedy przerwać doklejanie części."""
    job = CHAT_JOBS.get(rid or "")
    if job is None:
        return True
    waited = 0.0
    while job.get("paused") and not job.get("stop") and waited < 1800:
        time.sleep(0.4)
        waited += 0.4
    return not job.get("stop")

class ChatJobControl(BaseModel):
    action: str  # pause | resume | stop

@app.get("/chat-job/{rid}")
def chat_job_status_ep(rid: str, user: str = Depends(get_current_user)):
    job = CHAT_JOBS.get(rid)
    if job is None:
        raise HTTPException(status_code=404, detail="Zadanie czatu nie istnieje (zakończone albo wygasło).")
    return {"stage": job["stage"], "parts": job["parts"], "paused": bool(job.get("paused")),
            "stop": bool(job.get("stop")), "elapsed": int(time.time() - job["started"])}

@app.post("/chat-job/{rid}/control")
def chat_job_control_ep(rid: str, r: ChatJobControl, user: str = Depends(get_current_user)):
    job = CHAT_JOBS.get(rid)
    if job is None:
        raise HTTPException(status_code=404, detail="Zadanie czatu nie istnieje (zakończone albo wygasło).")
    action = (r.action or "").strip().lower()
    if action == "pause":
        job["paused"] = True
    elif action == "resume":
        job["paused"] = False
    elif action == "stop":
        job["stop"] = True
        job["paused"] = False
    else:
        raise HTTPException(status_code=400, detail=f"Nieznana akcja: {action}")
    return {"stage": job["stage"], "parts": job["parts"], "paused": bool(job.get("paused")),
            "stop": bool(job.get("stop")), "elapsed": int(time.time() - job["started"])}


# --- PLIK_W_TLE_V8: tworzenie dużego pliku jako zadanie w tle (karta postępu) ---
# Duży plik nie blokuje rozmowy: /chat zwraca od razu kartę zadania, a worker
# w osobnym wątku dopisuje kolejne części pliku (każda = osobne wywołanie LLM,
# rotacja kluczy jak w PLIK_NA_RATY_V7). Użytkownik może w trakcie dalej
# rozmawiać, a z karty: wstrzymać (pauza), wznowić (play), zatrzymać (stop –
# częściowy plik zostaje do pobrania) albo usunąć zadanie z kartą.
FILE_JOBS: dict = {}
FILE_JOB_MAX_PARTS = 24
FILE_JOB_MAX_CHARS = 400_000
FILE_JOBS_BOX_FOLDER = "filejobs"  # folder w Boxie użytkownika na zapisy postępu zadań

# --- Trwałość w Boxie: postęp zadania zapisuje się w chmurze użytkownika ---
# Dzięki temu po restarcie serwera albo po przywołaniu zapisanej rozmowy
# zadanie da się wznowić od miejsca wstrzymania, a nie od początku.
def _file_job_box_path(job_id: str) -> str:
    return f"{FILE_JOBS_BOX_FOLDER}/{job_id}.json"

def _box_upload_json(path_in_box: str, obj: dict) -> bool:
    try:
        data = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        files = {"file": (os.path.basename(path_in_box), data, "application/json")}
        r = requests.post(f"{BOX_BACKEND_URL}/upload", headers=_box_headers(),
                          files=files, data={"path": os.path.dirname(path_in_box)}, timeout=60)
        r.raise_for_status()
        return True
    except Exception as e:
        print(f"BOX save error ({path_in_box}): {e}")
        return False

def _box_download_json(path_in_box: str) -> dict | None:
    try:
        enc = requests.utils.quote(path_in_box, safe="")
        bt = _box_login()
        r = requests.get(f"{BOX_BACKEND_URL}/download/{enc}?token={bt}&preview=true", timeout=30)
        r.raise_for_status()
        data = r.json()
        if isinstance(data, dict):
            if "content" in data or "job_id" in data:
                return data
            raw = data.get("data")
            if isinstance(raw, str):
                try:
                    return json.loads(base64.b64decode(raw).decode("utf-8"))
                except Exception:
                    try:
                        return json.loads(raw)
                    except Exception:
                        return None
        return None
    except Exception as e:
        print(f"BOX load error ({path_in_box}): {e}")
        return None

def _box_delete_file(path_in_box: str) -> bool:
    try:
        resp = requests.post(f"{BOX_BACKEND_URL}/delete", headers=_box_headers(),
                             json={"path": path_in_box}, timeout=20)
        resp.raise_for_status()
        return True
    except Exception as e:
        print(f"BOX delete error ({path_in_box}): {e}")
        return False

def _file_job_save(job_id: str):
    """Zapisuje pełny postęp zadania (z treścią pliku) do Boxa. Nigdy nie wywraca workera."""
    job = FILE_JOBS.get(job_id)
    if job is None:
        return
    try:
        snapshot = {
            "job_id": job_id,
            "filename": job["filename"],
            "status": job["status"],
            "parts": job["parts"],
            "chars": job["chars"],
            "estimate": job["estimate"],
            "truncated": job["truncated"],
            "content": "".join(job["content_parts"]),
            "saved_at": datetime.utcnow().isoformat() + "Z",
        }
        if _box_upload_json(_file_job_box_path(job_id), snapshot):
            print(f"FILE_JOB {job_id}: postęp zapisany w Boxie ({job['chars']} znaków, część {job['parts']})")
    except Exception as e:
        print(f"FILE_JOB {job_id}: błąd zapisu do Boxa: {e}")

def _file_job_percent(job) -> int:
    if job["status"] == "done":
        return 100
    chars = job["chars"]
    if chars <= 0:
        return 0
    # Procent liczony od szacunku modelu ([SZACUJĘ:]); gdy go brak albo plik
    # urósł ponad szacunek, mianownik rośnie razem z plikiem (zostaje < 100%).
    eff = max(job["estimate"] or 0, int(chars * 1.25), 8000)
    return min(99, int(round(100 * chars / eff)))

def _file_job_snapshot(job_id: str) -> dict | None:
    job = FILE_JOBS.get(job_id)
    if not job:
        return None
    return {
        "job_id": job_id,
        "status": job["status"],
        "filename": job["filename"],
        "percent": _file_job_percent(job),
        "chars": job["chars"],
        "parts": job["parts"],
        "estimate": job["estimate"],
        "truncated": job["truncated"],
        "error": job["error"],
        "created": job.get("created"),
    }

def _file_job_continuation_instr(filename: str, full_tail: str) -> str:
    tail = full_tail[-220:].replace("\n", " ")
    return (
        "[SYSTEM] Twój plik " + filename + " został ucięty w połowie "
        "(limit jednej odpowiedzi). Kontynuuj automatycznie, bez pytania "
        "użytkownika i bez żadnego komentarza. Wyślij TYLKO dalszy ciąg treści "
        "pliku " + filename + ", dokładnie od miejsca, w którym urwałeś. "
        "Ostatnie znaki, które już wysłałeś: „…" + tail + "” – NIE powtarzaj ich "
        "ani niczego, co już jest w pliku. Nie dodawaj tagu [FILE:] ani "
        "znaczników ```. Gdy plik będzie kompletny w całości, zakończ go "
        "tagiem [/FILE] i niczym więcej."
    )

def _file_job_worker(job_id: str):
    global current_key_index
    job = FILE_JOBS.get(job_id)
    if job is None:
        return
    msgs = job["_msgs"]
    last_raw = job["_last_raw"]
    try:
        while True:
            if FILE_JOBS.get(job_id) is not job:
                return  # zadanie usunięte przez użytkownika
            if job["stop_requested"]:
                job["status"] = "stopped"
                job["truncated"] = True
                _file_job_save(job_id)
                return
            while job["status"] == "paused":
                time.sleep(0.4)
                if FILE_JOBS.get(job_id) is not job or job["stop_requested"]:
                    if job["stop_requested"]:
                        job["status"] = "stopped"
                        job["truncated"] = True
                        _file_job_save(job_id)
                    return
            if job["parts"] >= FILE_JOB_MAX_PARTS or job["chars"] >= FILE_JOB_MAX_CHARS:
                job["status"] = "done"
                job["truncated"] = True
                print(f"FILE_JOB {job_id}: limit części/rozmiaru – plik częściowy ({job['chars']} znaków)")
                _file_job_save(job_id)
                return
            full_tail = "".join(job["content_parts"])
            call_msgs = msgs + [
                {'role': 'model', 'parts': [last_raw]},
                {'role': 'user', 'parts': [_file_job_continuation_instr(job["filename"], full_tail)]},
            ]
            keys_try = ([job["api_key"]] if job["api_key"] else []) + api_keys[current_key_index:] + api_keys[:current_key_index]
            nxt = None
            for idx2, k2 in enumerate(keys_try):
                if idx2 > 15: break
                try:
                    nxt = call_llm(k2, call_msgs, SYSTEM_INSTRUCTION)
                    if k2 in api_keys: current_key_index = api_keys.index(k2)
                    break
                except Exception as e:
                    print(f"FILE_JOB key fail: {e}")
            if not nxt:
                job["status"] = "error"
                job["error"] = "Wszystkie klucze API zawiodły podczas tworzenia pliku."
                _file_job_save(job_id)
                return
            if job["stop_requested"] or FILE_JOBS.get(job_id) is not job:
                if FILE_JOBS.get(job_id) is job:
                    job["status"] = "stopped"
                    job["truncated"] = True
                    _file_job_save(job_id)
                return  # części z wywołania po stopie nie doklejamy
            cont = nxt.strip()
            cont = re.sub(r'^\[FILE:\s*[^\]]+?\]\s*', '', cont, flags=re.IGNORECASE)
            cont = re.sub(r'^```[\w]*\s*\n?', '', cont)
            closed = "[/FILE]" in cont.upper()
            if closed:
                cont = re.split(r'\[/FILE\]', cont, flags=re.IGNORECASE)[0]
                cont = re.sub(r'\n?```$', '', cont.rstrip())
            job["content_parts"].append(cont)
            job["chars"] += len(cont)
            job["parts"] += 1
            msgs = call_msgs
            last_raw = nxt
            # zapis kontekstu w zadaniu (żeby dało się je restartować po stopie)
            # oraz trwały zapis postępu w Boxie po każdej części
            job["_msgs"] = call_msgs
            job["_last_raw"] = nxt
            print(f"FILE_JOB {job_id}: część {job['parts']} (+{len(cont)} znaków, razem {job['chars']}, {_file_job_percent(job)}%)")
            if closed:
                job["status"] = "done"
                job["truncated"] = False
            _file_job_save(job_id)
            if closed:
                return
    except Exception as e:
        if FILE_JOBS.get(job_id) is job:
            job["status"] = "error"
            job["error"] = str(e)

def create_file_job(filename: str, first_content: str, base_msgs: list, last_raw: str,
                    api_key: str | None, estimate: int | None) -> str:
    """Startuje tworzenie pliku w wątku w tle. Zwraca job_id."""
    job_id = uuid.uuid4().hex[:12]
    FILE_JOBS[job_id] = {
        "status": "running",
        "filename": filename,
        "content_parts": [first_content] if first_content else [],
        "chars": len(first_content or ""),
        "parts": 1,
        "estimate": estimate,
        "truncated": False,
        "error": None,
        "stop_requested": False,
        "api_key": api_key,
        "created": time.time(),
        "_msgs": list(base_msgs),
        "_last_raw": last_raw,
    }
    # Wstępny zapis postępu w Boxie – w osobnym wątku, żeby nie opóźniać odpowiedzi czatu
    threading.Thread(target=_file_job_save, args=(job_id,), daemon=True).start()
    threading.Thread(target=_file_job_worker, args=(job_id,), daemon=True).start()
    print(f"FILE_JOB {job_id}: start tworzenia pliku {filename} w tle (szacunek: {estimate})")
    return job_id

class FileJobControl(BaseModel):
    action: str  # pause | resume | stop | delete

@app.get("/file-job/{job_id}")
def file_job_status_ep(job_id: str, user: str = Depends(get_current_user)):
    snap = _file_job_snapshot(job_id)
    if snap is None:
        raise HTTPException(status_code=404, detail="Zadanie nie istnieje (mógł wygasnąć przy restarcie serwera).")
    return snap

@app.post("/file-job/{job_id}/control")
def file_job_control_ep(job_id: str, r: FileJobControl, user: str = Depends(get_current_user)):
    job = FILE_JOBS.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Zadanie nie istnieje (mógł wygasnąć przy restarcie serwera).")
    action = (r.action or "").strip().lower()
    if action == "pause":
        if job["status"] == "running":
            job["status"] = "paused"
    elif action == "resume":
        if job["status"] == "paused":
            job["status"] = "running"
    elif action == "stop":
        job["stop_requested"] = True
        if job["status"] == "paused":
            job["status"] = "stopped"
            job["truncated"] = True
    elif action == "delete":
        job["stop_requested"] = True
        FILE_JOBS.pop(job_id, None)
        # Usuwamy też zapis postępu z Boxa, żeby nie dało się wznowić usuniętego zadania
        threading.Thread(target=_box_delete_file, args=(_file_job_box_path(job_id),), daemon=True).start()
        return {"status": "deleted", "job_id": job_id}
    else:
        raise HTTPException(status_code=400, detail=f"Nieznana akcja: {action}")
    if action in ("pause", "stop"):
        threading.Thread(target=_file_job_save, args=(job_id,), daemon=True).start()
    return _file_job_snapshot(job_id)

@app.get("/file-job/{job_id}/download")
def file_job_download_ep(job_id: str, user: str = Depends(get_current_user)):
    job = FILE_JOBS.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Zadanie nie istnieje (mógł wygasnąć przy restarcie serwera).")
    content = "".join(job["content_parts"])
    return make_file_response(job["filename"], content)

def _file_job_restore(job_id: str) -> dict | None:
    """Wznawia zadanie: z pamięci (pauza/stop) albo z zapisu w Boxie (po restarcie
    serwera / przywołaniu zapisanej rozmowy). Zwraca snapshot albo None, gdy nie
    ma ani zadania, ani zapisu w chmurze."""
    job = FILE_JOBS.get(job_id)
    if job is not None:
        if job["status"] == "paused":
            job["status"] = "running"
        elif job["status"] == "stopped":
            job["stop_requested"] = False
            job["status"] = "running"
            job["truncated"] = False
            threading.Thread(target=_file_job_worker, args=(job_id,), daemon=True).start()
        return _file_job_snapshot(job_id)

    data = _box_download_json(_file_job_box_path(job_id))
    if not data or "content" not in data:
        return None
    content = data.get("content") or ""
    filename = data.get("filename") or "plik.txt"
    tail = content[-6000:]
    FILE_JOBS[job_id] = {
        "status": "running",
        "filename": filename,
        "content_parts": [content] if content else [],
        "chars": len(content),
        "parts": int(data.get("parts") or 1),
        "estimate": data.get("estimate"),
        "truncated": False,
        "error": None,
        "stop_requested": False,
        "api_key": None,  # prywatny klucz użytkownika nie jest zapisywany – worker użyje puli
        "created": time.time(),
        # Odtworzony kontekst: model dostaje koniec pliku i polecenie kontynuacji,
        # więc dopisuje dalszy ciąg od miejsca wstrzymania, bez zaczynania od nowa.
        "_msgs": [{'role': 'user', 'parts': [
            f"[ODTWORZONE ZADANIE] Kontynuujesz tworzenie pliku {filename} po przerwie. "
            "Dotychczasowa treść jest już zapisana; dopisz TYLKO dalszy ciąg od miejsca "
            "urwania (koniec treści dostajesz w kolejnej wiadomości jako swoją poprzednią "
            "odpowiedź) i zamknij plik tagiem [/FILE], gdy będzie kompletny."]}],
        "_last_raw": f"[FILE: {filename}]\n{tail}",
    }
    print(f"FILE_JOB {job_id}: odtworzono z Boxa ({len(content)} znaków) – kontynuuję w tle")
    threading.Thread(target=_file_job_worker, args=(job_id,), daemon=True).start()
    return _file_job_snapshot(job_id)

@app.post("/file-job/{job_id}/resume")
def file_job_resume_ep(job_id: str, user: str = Depends(get_current_user)):
    snap = _file_job_restore(job_id)
    if snap is None:
        raise HTTPException(status_code=404, detail="Brak zapisanego postępu tego zadania w chmurze.")
    return snap


# --- KOMPILATOR INSTALATORÓW NSIS (wymaga pakietu systemowego 'nsis' na serwerze) ---
NSIS_INCLUDE_WHITELIST = {
    "MUI2.nsh", "MUI.nsh", "LogicLib.nsh", "FileFunc.nsh", "Sections.nsh",
    "WinMessages.nsh", "x64.nsh", "Memento.nsh", "nsDialogs.nsh",
    "InstallOptions.nsh", "LangFile.nsh", "Colors.nsh", "MultiUser.nsh",
}
NSIS_BANNED_DIRECTIVES = ("!system", "!execute", "!packhdr", "!appendfile", "!delfile")

def build_nsis_installer(exe_name: str, nsi_script: str, sources: dict) -> bytes:
    """
    Kompiluje instalator Windows (.exe) przez makensis.
    sources: {nazwa_pliku: bajty} - pliki, do których odwołuje się skrypt (z tagów [FILE:]).
    Zwraca bajty gotowego instalatora.
    """
    import shutil
    import subprocess
    import tempfile

    makensis = shutil.which("makensis")
    if not makensis:
        raise RuntimeError("makensis niedostępny na serwerze - zainstaluj pakiet systemowy 'nsis' "
                           "(packages.txt / Dockerfile), to NIE jest pakiet Pythona.")

    # Bezpieczeństwo: skrypt NSIS nie może wykonywać poleceń na serwerze w czasie kompilacji
    lowered = nsi_script.lower()
    for bad in NSIS_BANNED_DIRECTIVES:
        if bad in lowered:
            raise RuntimeError(f"Zabroniona dyrektywa w skrypcie NSIS: {bad}")
    for inc in re.findall(r'!include\s+"([^"]+)"', nsi_script, re.IGNORECASE):
        base = os.path.basename(inc)
        if "/" in inc or "\\" in inc or base not in NSIS_INCLUDE_WHITELIST:
            raise RuntimeError(f"Niedozwolony !include w skrypcie NSIS: {inc}")

    build_dir = tempfile.mkdtemp(prefix="nsis_build_")
    try:
        for fname, data in sources.items():
            safe = os.path.basename(fname)
            with open(os.path.join(build_dir, safe), "wb") as f:
                f.write(data)
        script_path = os.path.join(build_dir, "installer.nsi")
        with open(script_path, "w", encoding="utf-8") as f:
            f.write(nsi_script)

        result = subprocess.run(
            [makensis, "installer.nsi"],
            cwd=build_dir, capture_output=True, text=True, timeout=120,
        )
        if result.returncode != 0:
            tail = ((result.stdout or "") + "\n" + (result.stderr or ""))[-800:]
            raise RuntimeError(f"makensis zakończył się błędem: {tail}")

        exe_path = os.path.join(build_dir, exe_name)
        if not os.path.exists(exe_path):
            exes = [f for f in os.listdir(build_dir) if f.lower().endswith(".exe")]
            if not exes:
                raise RuntimeError("makensis nie wyprodukował pliku .exe")
            exe_path = os.path.join(build_dir, exes[0])
        with open(exe_path, "rb") as f:
            data = f.read()
        print(f"NSIS installer OK: {exe_name} ({len(data)} bytes)")
        return data
    finally:
        shutil.rmtree(build_dir, ignore_errors=True)



# --- AKUMULATOR_V3: budowanie na raty (oszczędzanie tokenów) ---
_AKUM_SKIP_SUFFIXES = ("_blad.txt", "_status.txt", "_info.txt", ".docx", ".pdf", ".exe", ".job")

def _akum_history_files(history):
    """Pliki projektu z poprzednich odpowiedzi asystenta. Frontend trzyma
    wygenerowane pliki jako załączniki wiadomości w historii – zbieramy je
    wszystkie, w kolejności rozmowy (przy łączeniu nowsza wersja pliku
    nadpisuje starszą)."""
    out = []
    if not history:
        return out
    for item in history:
        if getattr(item, "isHidden", False):
            continue
        if getattr(item, "role", None) not in ("model", "assistant"):
            continue
        for a in (getattr(item, "attachments", None) or []):
            name = (getattr(a, "name", "") or "").strip()
            data = getattr(a, "data", "") or ""
            if not name or not data:
                continue
            if name.lower().endswith(_AKUM_SKIP_SUFFIXES) or name.startswith("build_"):
                continue
            out.append({"filename": name,
                        "mime_type": getattr(a, "mime_type", None) or guess_mime(name),
                        "data": data})
    return out

def _akum_user_call_html(history):
    """Duży plik użytkownika do assets: cały call.html z załącznika albo
    złożony z części call-czesc-N-z-M.txt. Zwraca tekst albo None.
    Dzięki temu model nigdy nie musi przepisywać dużego pliku."""
    if not history:
        return None
    full = None
    parts = {}
    total = None
    for item in history:
        if getattr(item, "isHidden", False):
            continue
        if getattr(item, "role", None) != "user":
            continue
        for a in (getattr(item, "attachments", None) or []):
            name = (getattr(a, "name", "") or "").strip()
            data = getattr(a, "data", "") or ""
            if not name or not data:
                continue
            try:
                raw = base64.b64decode(data)
            except Exception:
                continue
            if name == "call.html":
                full = raw
            else:
                m = re.match(r"call-czesc-(\d+)-z-(\d+)\.txt$", name)
                if m:
                    parts[int(m.group(1))] = raw
                    total = int(m.group(2))
    if full is not None:
        try:
            return full.decode("utf-8")
        except Exception:
            return None
    if total and len(parts) == total and all(i in parts for i in range(1, total + 1)):
        try:
            return b"".join(parts[i] for i in range(1, total + 1)).decode("utf-8")
        except Exception:
            return None
    return None

def _akum_is_asset_attachment(name) -> bool:
    """AKUMULATOR_V4 (OSZCZEDNY): czy załącznik to duży plik do buildu
    (call.html albo jego części)? Takie pliki nie idą do modelu."""
    n = (name or "").strip()
    return n == "call.html" or re.match(r"call-czesc-\d+-z-\d+\.txt$", n) is not None

def _trim_file_blocks(text: str) -> str:
    """W historii wysyłanej do modelu długie bloki [FILE:] z poprzednich
    odpowiedzi zastępujemy krótką notką – treść plików jest zapisana
    w załącznikach rozmowy, a bloki zjadały większość tokenów wejścia."""
    def _sub(m):
        if len(m.group(2)) <= 400:
            return m.group(0)
        return (f"[FILE: {m.group(1).strip()}]\n"
                "[treść pliku pominięta w historii dla oszczędności tokenów – "
                "plik jest zapisany w rozmowie jako załącznik]\n[/FILE]")
    text = re.sub(r"\[FILE:\s*([^\]]+?)\]\s*(.*?)\s*\[/FILE\]", _sub, text,
                  flags=re.DOTALL | re.IGNORECASE)
    # niedomknięty [FILE:] na końcu starej odpowiedzi (była ucięta) – też skracamy,
    # inaczej jeden taki blok zje większość tokenów wejścia kolejnej odpowiedzi
    def _sub_unclosed(m):
        if len(m.group(2)) <= 400:
            return m.group(0)
        return (f"[FILE: {m.group(1).strip()}]\n"
                "[treść pliku pominięta w historii dla oszczędności tokenów – "
                "plik był ucięty w oryginalnej odpowiedzi]\n[/FILE]")
    return re.sub(r"\[FILE:\s*([^\]]+?)\]\s*(.*)$", _sub_unclosed, text,
                  flags=re.DOTALL | re.IGNORECASE)

def _extract_and_build_files(ai_text: str, history=None) -> tuple[str, list]:
    """
    Wyciąga z odpowiedzi AI tagi [FILE:], [GENERATE_DOCX:], [GENERATE_PDF:], [NSIS_INSTALLER:]
    i buduje listę plików do zwrócenia.
    Zwraca (czysty_tekst_bez_tagów, lista_plików).
    """
    files = []
    clean = ai_text
    truncated_files = []  # nazwy plików z niedomkniętych tagów (odpowiedź ucięta przez limit tokenów)

    # 1. [FILE: name.ext] ... [/FILE]
    file_pattern = re.compile(
        r'\[FILE:\s*([^\]]+?)\]\s*(.*?)\s*\[/FILE\]',
        re.DOTALL | re.IGNORECASE
    )
    for m in file_pattern.finditer(ai_text):
        fname = m.group(1).strip()
        content = m.group(2)
        # usuń ewentualne otaczające ```language
        content = re.sub(r'^```[\w]*\n?', '', content.strip())
        content = re.sub(r'\n?```$', '', content)
        files.append(make_file_response(fname, content))
    clean = file_pattern.sub('', clean)

    # 1b. NIEDOMKNIĘTY [FILE: name.ext] bez [/FILE] – odpowiedź modelu urwała się
    #     w połowie pliku (limit tokenów wyjścia, np. długa gra HTML).
    #     Bez tego cała treść pliku lądowała w czacie jako "kod zamiast pliku".
    #     Zasada: wszystko od tagu do końca odpowiedzi to treść pliku – wyciągamy
    #     ją do pobrania tak samo jak domknięty plik i dopisujemy ostrzeżenie.
    unclosed_file_pattern = re.compile(
        r'\[FILE:\s*([^\]]+?)\]\s*(.*)$',
        re.DOTALL | re.IGNORECASE
    )
    m_unclosed = unclosed_file_pattern.search(clean)
    if m_unclosed:
        fname = m_unclosed.group(1).strip()
        content = m_unclosed.group(2)
        content = re.sub(r'^```[\w]*\n?', '', content.strip())
        content = re.sub(r'\n?```$', '', content)
        if content.strip():
            files.append(make_file_response(fname, content))
            truncated_files.append(fname)
            print(f"FILE_TRUNCATED: wyciągnięto niedomknięty plik {fname} ({len(content)} znaków)")
        clean = unclosed_file_pattern.sub('', clean)

    # 2. [GENERATE_DOCX: Tytuł] ... [/GENERATE_DOCX]
    docx_pattern = re.compile(
        r'\[GENERATE_DOCX:\s*([^\]]*)\]\s*(.*?)\s*\[/GENERATE_DOCX\]',
        re.DOTALL | re.IGNORECASE
    )
    for m in docx_pattern.finditer(ai_text):
        title = m.group(1).strip() or "Dokument"
        body = m.group(2).strip()
        try:
            docx_bytes = create_docx_from_text(title, body)
            safe_name = re.sub(r'[^\w\s\-ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]', '', title)[:40].strip() or "dokument"
            safe_name = safe_name.replace(' ', '_') + ".docx"
            files.append(make_file_response(safe_name, docx_bytes))
        except Exception as e:
            print(f"DOCX gen error: {e}")
    clean = docx_pattern.sub('', clean)

    # 3. [GENERATE_PDF: Tytuł] ... [/GENERATE_PDF]
    pdf_pattern = re.compile(
        r'\[GENERATE_PDF:\s*([^\]]*)\]\s*(.*?)\s*\[/GENERATE_PDF\]',
        re.DOTALL | re.IGNORECASE
    )
    for m in pdf_pattern.finditer(ai_text):
        title = m.group(1).strip() or "Dokument"
        body = m.group(2).strip()
        try:
            pdf_bytes = create_pdf_from_text(title, body)
            if not pdf_bytes or len(pdf_bytes) < 50:
                raise RuntimeError("PDF output too small / empty")
            safe_name = re.sub(r'[^\w\s\-ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]', '', title)[:40].strip() or "dokument"
            safe_name = safe_name.replace(' ', '_') + ".pdf"
            files.append(make_file_response(safe_name, pdf_bytes, "application/pdf"))
            print(f"PDF generated OK: {safe_name} ({len(pdf_bytes)} bytes)")
        except Exception as e:
            print(f"PDF gen error: {e}")
            # awaryjnie oddaj treść jako .txt, żeby użytkownik coś dostał
            safe_name = re.sub(r'[^\w\s\-ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]', '', title)[:40].strip() or "dokument"
            safe_name = safe_name.replace(' ', '_') + ".txt"
            files.append(make_file_response(safe_name, f"{title}\n\n{body}"))
    clean = pdf_pattern.sub('', clean)

    # 3b. NIEDOMKNIĘTE [GENERATE_DOCX:] / [GENERATE_PDF:] bez tagów zamykających
    #     (odpowiedź ucięta) – generujemy dokument z tego, co doszło do końca tekstu.
    unclosed_docx = re.search(r'\[GENERATE_DOCX:\s*([^\]]*)\]\s*(.*)$', clean, re.DOTALL | re.IGNORECASE)
    if unclosed_docx:
        title = unclosed_docx.group(1).strip() or "Dokument"
        body = unclosed_docx.group(2).strip()
        if body:
            try:
                docx_bytes = create_docx_from_text(title, body)
                safe_name = re.sub(r'[^\w\s\-ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]', '', title)[:40].strip() or "dokument"
                files.append(make_file_response(safe_name.replace(' ', '_') + ".docx", docx_bytes))
                truncated_files.append(safe_name.replace(' ', '_') + ".docx")
            except Exception as e:
                print(f"DOCX gen error (unclosed): {e}")
        clean = re.sub(r'\[GENERATE_DOCX:\s*[^\]]*\]\s*.*$', '', clean, flags=re.DOTALL | re.IGNORECASE)

    unclosed_pdf = re.search(r'\[GENERATE_PDF:\s*([^\]]*)\]\s*(.*)$', clean, re.DOTALL | re.IGNORECASE)
    if unclosed_pdf:
        title = unclosed_pdf.group(1).strip() or "Dokument"
        body = unclosed_pdf.group(2).strip()
        if body:
            try:
                pdf_bytes = create_pdf_from_text(title, body)
                safe_name = re.sub(r'[^\w\s\-ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]', '', title)[:40].strip() or "dokument"
                files.append(make_file_response(safe_name.replace(' ', '_') + ".pdf", pdf_bytes, "application/pdf"))
                truncated_files.append(safe_name.replace(' ', '_') + ".pdf")
            except Exception as e:
                print(f"PDF gen error (unclosed): {e}")
                safe_name = re.sub(r'[^\w\s\-ąćęłńóśźżĄĆĘŁŃÓŚŹŻ]', '', title)[:40].strip() or "dokument"
                files.append(make_file_response(safe_name.replace(' ', '_') + ".txt", f"{title}\n\n{body}"))
        clean = re.sub(r'\[GENERATE_PDF:\s*[^\]]*\]\s*.*$', '', clean, flags=re.DOTALL | re.IGNORECASE)

    # 4. [NSIS_INSTALLER: Nazwa-Setup.exe] ...skrypt .nsi... [/NSIS_INSTALLER]
    #    Kompiluje prawdziwy instalator Windows z plików [FILE:] z tej samej odpowiedzi.
    nsis_pattern = re.compile(
        r'\[NSIS_INSTALLER:\s*([^\]]+?)\]\s*(.*?)\s*\[/NSIS_INSTALLER\]',
        re.DOTALL | re.IGNORECASE
    )
    nsis_matches = list(nsis_pattern.finditer(ai_text))
    if nsis_matches:
        sources = {}
        for f in files:
            try:
                sources[f["filename"]] = base64.b64decode(f["data"])
            except Exception:
                continue
        for m in nsis_matches:
            exe_name = os.path.basename(m.group(1).strip()) or "Setup.exe"
            if not exe_name.lower().endswith(".exe"):
                exe_name += ".exe"
            script = m.group(2)
            script = re.sub(r'^```[\w]*\n?', '', script.strip())
            script = re.sub(r'\n?```$', '', script)
            try:
                exe_bytes = build_nsis_installer(exe_name, script, sources)
                files.append(make_file_response(exe_name, exe_bytes, "application/octet-stream"))
            except Exception as e:
                print(f"NSIS build error: {e}")
                files.append(make_file_response(
                    exe_name.replace(".exe", "_blad.txt"),
                    f"Nie udało się skompilować instalatora {exe_name}:\n{e}"
                ))
    clean = nsis_pattern.sub('', clean)

    # 5. [GH_BUILD: target] ... [/GH_BUILD] – build w chmurze (GitHub Actions)
    #    Backend NIE buduje synchronicznie: zwraca kartę buildu (pseudo-plik
    #    application/x-build-job z projektem w środku), z której frontend
    #    startuje build w tle przez POST /cloud-build/start i pokazuje
    #    status na żywo. Pliki projektu znikają z listy plików do pobrania
    #    (są osiągalne z karty przyciskiem "Pobierz sam projekt").
    gh_pattern = re.compile(
        r'\[GH_BUILD:\s*([^\]]+?)\]\s*(.*?)\s*\[/GH_BUILD\]',
        re.DOTALL | re.IGNORECASE
    )
    gh_matches = list(gh_pattern.finditer(ai_text))
    if gh_matches:
        # AKUMULATOR_V3: źródła buildu = pliki z całej rozmowy (załączniki
        # poprzednich odpowiedzi asystenta) + pliki z bieżącej odpowiedzi.
        # Na końcu dokładamy duży plik użytkownika z jego załączników –
        # nadpisuje ewentualną wersję od modelu, bo jest bajt w bajt.
        merged = {}
        for hf in _akum_history_files(history):
            merged[hf["filename"]] = hf
        for f in files:
            if not f["filename"].endswith(("_blad.txt", "_status.txt", "_info.txt",
                                           ".docx", ".pdf", ".exe")):
                merged[f["filename"]] = f
        if any(m.group(1).strip().lower() == "android" for m in gh_matches):
            asset_html = _akum_user_call_html(history)
            if asset_html is not None:
                merged["app/src/main/assets/call.html"] = make_file_response(
                    "app/src/main/assets/call.html", asset_html)
                print(f"AKUMULATOR_V3: assets/call.html z załączników użytkownika "
                      f"({len(asset_html)} znaków)")
        gh_src_files = list(merged.values())
        if gh_src_files:
            src_names = {f["filename"] for f in gh_src_files}
            files = [f for f in files if f["filename"] not in src_names]
            for m in gh_matches:
                target = m.group(1).strip().lower()
                payload = {"target": target, "sources": gh_src_files}
                files.append(make_file_response(
                    f"build_{target}.job",
                    json.dumps(payload, ensure_ascii=False),
                    "application/x-build-job",
                ))
        else:
            files.append(make_file_response(
                "build_blad.txt",
                "Nie udało się przygotować buildu w chmurze: brak plików projektu "
                "w odpowiedzi (dołącz je tagami [FILE:])."
            ))
    clean = gh_pattern.sub('', clean)

    # 6. [GH_BUILD_STATUS: numer] – dokończenie wcześniejszego builda z chmury
    ghs_pattern = re.compile(r'\[GH_BUILD_STATUS:\s*(\d+)\s*\]', re.IGNORECASE)
    for m in ghs_pattern.finditer(ai_text):
        run_id = int(m.group(1))
        try:
            outputs, info = github_build_status(run_id)
            for oname, odata in outputs.items():
                files.append(make_file_response(oname, odata, "application/octet-stream"))
            files.append(make_file_response(f"build_{run_id}_status.txt", info))
        except Exception as e:
            print(f"GH build status error: {e}")
            files.append(make_file_response(
                f"build_{run_id}_blad.txt",
                f"Nie udało się pobrać builda {run_id}:\n{e}"
            ))
    clean = ghs_pattern.sub('', clean)

    # 7. BLOKI KODU ``` W TEKŚCIE -> PLIKI DO POBRANIA (twarda zasada użytkownika:
    #    "nigdy kodu w czacie, tylko gotowe pliki"). Model mimo instrukcji czasem
    #    wkleja kod blokiem ``` zamiast tagiem [FILE:] – zamieniamy więc każdy
    #    większy blok (>= 200 znaków) na plik do pobrania i zdejmujemy go z tekstu.
    #    Krótkie fragmenty (komendy, przykłady 1-2 linijki) zostają w tekście.
    _LANG_EXT = {
        "python": "py", "py": "py", "javascript": "js", "js": "js", "typescript": "ts",
        "ts": "ts", "html": "html", "htm": "html", "css": "css", "json": "json",
        "markdown": "md", "md": "md", "sql": "sql", "bash": "sh", "shell": "sh",
        "sh": "sh", "java": "java", "c": "c", "cpp": "cpp", "c++": "cpp", "rust": "rs",
        "go": "go", "php": "php", "ruby": "rb", "yaml": "yaml", "yml": "yml",
        "xml": "xml", "plaintext": "txt", "text": "txt", "txt": "txt",
    }
    fence_pattern = re.compile(r'```(\w*)\n?(.*?)```', re.DOTALL)
    used_names = {f["filename"] for f in files}
    converted_names = []
    _fence_counter = [0]

    def _fence_to_file(m):
        lang = (m.group(1) or "").lower()
        code = (m.group(2) or "").strip("\n")
        if len(code.strip()) < 200:
            return m.group(0)  # krótki fragment – zostaje w tekście
        ext = _LANG_EXT.get(lang, lang if lang.isalnum() and len(lang) <= 5 else "txt")
        _fence_counter[0] += 1
        if ext == "html":
            base = "index.html" if _fence_counter[0] == 1 and "index.html" not in used_names else f"strona_{_fence_counter[0]}.html"
        else:
            base = f"kod_{_fence_counter[0]}.{ext}"
        name = base
        n = 2
        while name in used_names:
            stem, dot, e = base.rpartition(".")
            name = f"{stem}_{n}.{e}" if dot else f"{base}_{n}"
            n += 1
        used_names.add(name)
        files.append(make_file_response(name, code))
        converted_names.append(name)
        print(f"CODE_FENCE_TO_FILE: blok ```{lang} ({len(code)} znaków) -> plik {name}")
        return ""

    clean = fence_pattern.sub(_fence_to_file, clean)
    if converted_names:
        clean = clean.rstrip() + "\n\n📁 Kod zapisano jako plik do pobrania: " + ", ".join(converted_names) + "."

    # Ostrzeżenie o plikach z uciętej odpowiedzi (niedomknięte tagi)
    if truncated_files:
        clean = clean.rstrip() + (
            "\n\n⚠️ Uwaga: moja odpowiedź urwała się w trakcie tworzenia pliku: "
            + ", ".join(truncated_files)
            + " – plik może być niekompletny. Napisz „dokończ”, a dopiszę brakującą resztę."
        )

    # posprzątaj wielokrotne puste linie
    clean = re.sub(r'\n{3,}', '\n\n', clean).strip()
    return clean, files


@app.post("/chat")
def chat(r: ChatRequest, resp: Response, user: str = Depends(get_current_user)):
    global current_key_index
    _chat_job_register(r.request_id)  # chmurka oczekiwania może tym zadaniem sterować (pauza/stop)
    prompt_parts = []
    for i in [x for x in r.history if not x.isHidden]:
        _content = i.content
        if _content and i.role in ("model", "assistant"):
            _content = _trim_file_blocks(_content)  # AKUMULATOR_V3: odchudzanie historii
        parts = [_content] if _content else []
        if i.attachments:
            for a in i.attachments:
                # AKUMULATOR_V4 (OSZCZEDNY): części call.html i sam call.html NIE idą
                # do modelu – backend dokleja je do buildu z surowych załączników
                # (_akum_user_call_html), więc model dostaje tylko krótką notkę.
                # Bez tego jedno zapytanie ważyło ~33 tys. tokenów i darmowe
                # klucze (np. Groq, limit 8000/min) odrzucały je w całości (413).
                if _akum_is_asset_attachment(getattr(a, "name", "")):
                    parts.append(f"[załącznik: {a.name} – treść pominięta dla oszczędności tokenów; "
                                 "backend dołączy ten plik do buildu co do bajta]")
                else:
                    parts.append(process_att(a))
        prompt_parts.append({'role': i.role, 'parts': parts})

    loops = 0
    while loops < 3:
        loops += 1
        keys = ([r.apiKey] if r.apiKey else []) + api_keys[current_key_index:] + api_keys[:current_key_index]
        ai_text = None
        for idx, k in enumerate(keys):
            if idx > 15: break
            try:
                ai_text = call_llm(k, prompt_parts, SYSTEM_INSTRUCTION)
                if k in api_keys: current_key_index = api_keys.index(k)
                break
            except Exception as e: 
                print(f"Key fail: {e}")
        
        if not ai_text: 
            resp.status_code = 429
            return {"error": "All keys failed", "all_keys_exhausted": True}

        # --- najpierw narzędzia wymagające pętli (search/scrape) ---
        if "[TOOL:search]" in ai_text:
            q = re.search(r"\[TOOL:search\](.*?)\[/TOOL\]", ai_text, re.DOTALL).group(1)
            prompt_parts.append({'role': 'model', 'parts': [ai_text]})
            prompt_parts.append({'role': 'user', 'parts': [f"WYNIK: {search_web(q)}"]})
            continue
        if "[TOOL:scrape]" in ai_text:
            u = re.search(r"\[TOOL:scrape\](.*?)\[/TOOL\]", ai_text, re.DOTALL).group(1)
            prompt_parts.append({'role': 'model', 'parts': [ai_text]})
            prompt_parts.append({'role': 'user', 'parts': [f"TREŚĆ: {scrape_web(u)}"]})
            continue

        # --- narzędzia obrazkowe / URL / STORY_PDF (zwracają od razu) ---
        user_msg = ai_text.split('[')[0].strip() if '[' in ai_text else ai_text.strip()

        if "[OPEN_URL:" in ai_text:
            try:
                u = re.search(r'url="([^"]*)"', ai_text).group(1)
                d = open_url_and_get_content(u)
                return {f"{d['type']}_file": {"filename":d['filename'], "data":d['data']}, "generated_text": user_msg or "Otwarto."}
            except: return {"generated_text": "Błąd URL."}
        elif "[GENERATE_IMAGE:" in ai_text:
            try:
                p = re.search(r'prompt="([^"]*)"', ai_text).group(1)
                return {"image_file": {"filename":"gen.png", "data": base64.b64encode(generate_image_from_prompt(p)).decode()}, "generated_text": user_msg or "Obraz:"}
            except: return {"generated_text": "Błąd gen."}
        elif "[EDIT_IMAGE:" in ai_text:
            try:
                att = next((a for i in reversed(r.history) if i.attachments for a in i.attachments if "image" in a.mime_type), None)
                if not att: return {"generated_text": "Brak obrazu."}
                return {"image_file": {"filename":"edit.png", "data": base64.b64encode(edit_image(att.data, ai_text)).decode()}, "generated_text": user_msg or "Zrobione."}
            except: return {"generated_text": "Błąd edycji."}
        elif "[CHANGE_BACKGROUND:" in ai_text:
            try:
                att = next((a for i in reversed(r.history) if i.attachments for a in i.attachments if "image" in a.mime_type), None)
                if not att: return {"generated_text": "Brak obrazu."}
                fg = remove(base64.b64decode(att.data))
                bg = generate_image_from_prompt(re.search(r'background_prompt="([^"]*)"', ai_text).group(1))
                return {"image_file": {"filename":"bg.png", "data": base64.b64encode(composite_images(fg, bg)).decode()}, "generated_text": user_msg or "Tło zmienione."}
            except: return {"generated_text": "Błąd tła."}
        elif "[INPAINT_IMAGE:" in ai_text:
            try:
                att = next((a for i in reversed(r.history) if i.attachments for a in i.attachments if "image" in a.mime_type), None)
                return {"image_file": {"filename":"inp.png", "data": base64.b64encode(inpaint_image(att.data, re.search(r'prompt="([^"]*)"', ai_text).group(1))).decode()}, "generated_text": user_msg or "Usunięto."}
            except: return {"generated_text": "Błąd inpaint."}
        elif "[CREATE_COMPOSITION:" in ai_text:
            try:
                m = re.search(r"\[CREATE_COMPOSITION:(.*?)\]", ai_text, re.DOTALL).group(1)
                def gp(k): return re.search(f'{k}="([^"]*)"', m).group(1)
                fg_b = base64.b64decode(find_attachment_in_history(gp("foreground_value"), r.history).data) if gp("foreground_source_type")=="file" else generate_image_from_prompt(gp("foreground_value"))
                bg_b = base64.b64decode(find_attachment_in_history(gp("background_value"), r.history).data) if gp("background_source_type")=="file" else generate_image_from_prompt(gp("background_value"))
                return {"image_file": {"filename":"comp.png", "data": base64.b64encode(composite_images(fg_b, bg_b)).decode()}, "generated_text": user_msg or "Kompozycja."}
            except: return {"generated_text": "Błąd komp."}
        elif "[STORY_PDF]" in ai_text:
            try:
                j = json.loads(re.search(r"\[STORY_PDF\](.*)\[/STORY_PDF\]", ai_text, re.DOTALL).group(1))
                return {"pdf_file": {"filename":"story.pdf", "data": base64.b64encode(create_story_pdf(j, r.history)).decode()}, "generated_text": user_msg or "PDF."}
            except: return {"generated_text": "Błąd PDF."}

        # --- AUTOBUILD_V6 + PLIK_NA_RATY_V7: dokańczanie bez zatrzymywania się ---
        # Backend sam dopytuje model w ramach TEGO zapytania, zamiast zmuszać
        # użytkownika do poganiania agenta. Każda kontynuacja to osobne wywołanie
        # LLM (osobny limit tokenów na część), a klucze rotują przy błędach
        # i między częściami – jedno zadanie może więc zużyć wielokrotnie więcej
        # tokenów, niż mieści jedna odpowiedź, i trwać tak długo, aż plik będzie
        # kompletny albo użytkownik przerwie czekanie przyciskiem stop (■).
        # Dwa przypadki:
        #  A) projekt (pliki .gradle/.csproj/.cpp/...) bez tagu buildu – jak dawniej;
        #  B) plik [FILE:] rozpoczęty, ale NIEDOMKNIĘTY (odpowiedź ucięta na limicie
        #     tokenów) – model dopisuje dalszy ciąg dokładnie od miejsca urwania,
        #     aż zamknie [/FILE]. Gdyby częściowy wynik i tak wrócił do użytkownika
        #     (limit części/rozmiaru), parser wyciągnie go jako plik z ostrzeżeniem.
        # Zabezpieczenia: maks. 24 kontynuacje i maks. ~400 tys. znaków całości.
        ai_full = ai_text
        _file_tag_re = re.compile(r"\[FILE:\s*([^\]]+?)\]", re.IGNORECASE)
        _file_close_re = re.compile(r"\[/FILE\]", re.IGNORECASE)
        def _pliki_projektu(txt):
            nazwy = {m.group(1).strip() for m in _file_tag_re.finditer(txt)}
            def _projektowa(n):
                return (n.endswith((".gradle", ".csproj", ".sln", ".cpp", ".h", ".hpp", ".java",
                                    ".kt", ".nsi", "CMakeLists.txt", "AndroidManifest.xml")))
            return nazwy, any(_projektowa(n) for n in nazwy)
        def _niedomkniety_plik(txt):
            opens = [m.group(1).strip() for m in _file_tag_re.finditer(txt)]
            if len(opens) > len(_file_close_re.findall(txt)):
                return opens[-1]
            return None

        # --- PLIK_W_TLE_V8: duży plik (niedomknięty, niebędący projektem do buildu)
        #     leci do zadania w tle z kartą postępu, a rozmowa wraca do użytkownika
        #     od razu – może dalej pisać, a plik dokończy się w tle. ---
        _niedomkniety_start = _niedomkniety_plik(ai_full)
        _nazwy_start, _czy_projekt_start = _pliki_projektu(ai_text)
        if (_niedomkniety_start and not _czy_projekt_start
                and "[GH_BUILD" not in ai_full and "[NSIS_INSTALLER" not in ai_full):
            _opens = list(_file_tag_re.finditer(ai_full))
            _m_open = _opens[-1]
            _first_content = ai_full[_m_open.end():]
            _first_content = re.sub(r'^\s*```[\w]*\n?', '', _first_content)
            _intro = re.sub(r'\[SZACUJĘ:[^\]]*\]', '', ai_full[:_m_open.start()]).strip()
            _est_m = re.search(r"\[SZACUJĘ:\s*~?\s*(\d+)", ai_full)
            _estimate = int(_est_m.group(1)) if _est_m else None
            _job_id = create_file_job(
                filename=_niedomkniety_start,
                first_content=_first_content,
                base_msgs=prompt_parts,
                last_raw=ai_text,
                api_key=r.apiKey,
                estimate=_estimate,
            )
            _job_payload = {"job_id": _job_id, "filename": _niedomkniety_start, "estimate": _estimate}
            return {
                "generated_text": _intro or (
                    f"Tworzę plik {_niedomkniety_start} w tle – postęp widać na karcie "
                    "poniżej. Możesz dalej rozmawiać, a plik dokończy się sam."
                ),
                "files": [make_file_response(
                    f"filejob_{_job_id}.job",
                    json.dumps(_job_payload, ensure_ascii=False),
                    "application/x-file-job",
                )],
            }

        _MAX_CONT = 24
        _MAX_FULL_CHARS = 400_000
        _cont = 0
        _chat_stopped = False
        while _cont < _MAX_CONT:
            if "[GH_BUILD" in ai_full or "[NSIS_INSTALLER" in ai_full:
                break
            if len(ai_full) >= _MAX_FULL_CHARS:
                print("PLIK_NA_RATY_V7: limit rozmiaru całości – koniec kontynuacji")
                break
            _niedomkniety = _niedomkniety_plik(ai_full)
            _nazwy, _czy_projekt = _pliki_projektu(ai_text)  # tylko najnowsza partia: sam tekst = model skonczyl
            if not _niedomkniety and not _czy_projekt:
                break
            # Sterowanie z chmurki: pauza czeka między częściami, stop przerywa pętlę
            _chat_job_set_stage(
                r.request_id,
                (f"Dopisuję dalszy ciąg pliku {_niedomkniety} (część {_cont + 2})…" if _niedomkniety
                 else "Zbieram kolejne pliki projektu…"),
                parts=_cont + 1)
            if not _chat_job_gate(r.request_id):
                _chat_stopped = True
                print("CHAT_JOB: użytkownik zatrzymał zadanie z chmurki – przerywam doklejanie")
                break
            prompt_parts.append({'role': 'model', 'parts': [ai_text]})
            if _niedomkniety:
                _tail = ai_full[-220:].replace("\n", " ")
                _instr = (
                    "[SYSTEM] Twój plik " + _niedomkniety + " został ucięty w połowie "
                    "(limit jednej odpowiedzi). Kontynuuj automatycznie, bez pytania "
                    "użytkownika i bez żadnego komentarza. Wyślij TYLKO dalszy ciąg treści "
                    "pliku " + _niedomkniety + ", dokładnie od miejsca, w którym urwałeś. "
                    "Ostatnie znaki, które już wysłałeś: „…" + _tail + "” – NIE powtarzaj ich "
                    "ani niczego, co już jest w pliku. Nie dodawaj tagu [FILE:] ani "
                    "znaczników ```. Gdy plik będzie kompletny w całości, zakończ go "
                    "tagiem [/FILE] i niczym więcej."
                )
            else:
                _instr = (
                    "[SYSTEM] Kontynuuj automatycznie, bez zatrzymywania się i bez pytania "
                    "użytkownika. Jeśli projekt jest kompletny, wyślij teraz TYLKO tag "
                    "zlecenia buildu (np. [GH_BUILD: android] [/GH_BUILD]). Jeśli brakuje "
                    "plików, wyślij kolejne 1–2 pliki [FILE:] tego samego projektu. "
                    "Bez podsumowań i bez list postępu."
                )
            prompt_parts.append({'role': 'user', 'parts': [_instr]})
            # Rotacja kluczy: każda część zaczyna próby od bieżącego klucza w puli,
            # a przy błędzie/limicie przechodzi na kolejne (jak w głównej pętli).
            _keys_try = ([r.apiKey] if r.apiKey else []) + api_keys[current_key_index:] + api_keys[:current_key_index]
            _nxt = None
            for _idx2, _k2 in enumerate(_keys_try):
                if _idx2 > 15: break
                try:
                    _nxt = call_llm(_k2, prompt_parts, SYSTEM_INSTRUCTION)
                    if _k2 in api_keys: current_key_index = api_keys.index(_k2)
                    break
                except Exception as _e:
                    print(f"Key fail: {_e}")
            if not _nxt:
                break
            if _niedomkniety:
                # Kontynuacja pliku: doklejamy BEZ separatora (urwanie bywa w pół
                # słowa), po oczyszczeniu początku z ewentualnego tagu [FILE:]
                # albo ogrodzenia ```, gdyby model dodał je mimo zakazu.
                _cont_txt = _nxt.strip()
                _cont_txt = re.sub(r'^\[FILE:\s*[^\]]+?\]\s*', '', _cont_txt, flags=re.IGNORECASE)
                _cont_txt = re.sub(r'^```[\w]*\s*\n?', '', _cont_txt)
                ai_text = _cont_txt
                ai_full += _cont_txt
                print(f"PLIK_NA_RATY_V7: doklejono część {_cont + 1} pliku {_niedomkniety} "
                      f"(+{len(_cont_txt)} znaków, razem {len(ai_full)})")
            else:
                ai_text = _nxt
                ai_full += "\n" + ai_text
                print(f"AUTOBUILD_V6: automatyczna kontynuacja nr {_cont + 1}")
            _cont += 1
        # awaryjnie: kompletny projekt Androida bez tagu -> dopisujemy zlecenie sami
        _nazwy_all, _ = _pliki_projektu(ai_full)
        if "[GH_BUILD" not in ai_full and any("AndroidManifest.xml" in n for n in _nazwy_all):
            ai_full += "\n[GH_BUILD: android]\n[/GH_BUILD]"
            print("AUTOBUILD_V6: dopisano brakujący tag [GH_BUILD: android]")

        # --- NOWE: pliki kodu / docx / pdf z tagów ---
        clean_text, generated_files = _extract_and_build_files(ai_full, r.history)
        if _chat_stopped:
            clean_text = ((clean_text or "") +
                          "\n\n⏹ Zatrzymałeś to zadanie z chmurki, więc wynik powyżej jest częściowy. "
                          "Napisz „dokończ”, a dopiszę resztę.").strip()

        result = {"generated_text": clean_text or user_msg}

        if generated_files:
            # kompatybilność wsteczna + nowa lista
            result["files"] = generated_files
            # jeśli jest dokładnie jeden plik obrazu/pdf – zostaw stare klucze
            if len(generated_files) == 1:
                f = generated_files[0]
                if f["mime_type"].startswith("image/"):
                    result["image_file"] = {"filename": f["filename"], "data": f["data"]}
                elif f["mime_type"] == "application/pdf":
                    result["pdf_file"] = {"filename": f["filename"], "data": f["data"]}

        return result
    
    return {"generated_text": "Limit operacji."}
